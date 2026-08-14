"""
Artefatos da comparacao entre planilhas (RF-REC-001).

Quatro saidas, cada uma para um leitor diferente:

- ``comparacao_report.json`` — para integracao e para o proximo passo do
  fluxo (o formato tem tudo: chave, colunas comparadas, normalizacoes,
  contadores, divergencias, conflitos, avisos e as LIMITACOES declaradas).
- ``divergencias.xlsx``      — para a pessoa que vai conferir: uma aba por
  categoria, com a formatacao dos demais artefatos do projeto.
- ``somente_a.csv`` / ``somente_b.csv`` — as linhas inteiras que existem em
  apenas uma das fontes, prontas para virar trabalho (cadastrar, remover).

Regra que atravessa o modulo: os valores gravados sao os ORIGINAIS das
fontes. A normalizacao serviu para PAREAR e para COMPARAR; o que o
operador ve e o que esta na planilha dele.

Nenhuma funcao aqui toca os arquivos de entrada.
"""

from __future__ import annotations

import json
from typing import TYPE_CHECKING, Any

import pandas as pd
from openpyxl import Workbook

from autotarefas.reconcile.result import (
    Category,
    ComparisonResult,
    KeyConflict,
    RecordComparison,
    TableSource,
)
from autotarefas.tasks.xlsx_style import (
    DATA_FONT,
    ERROR_FILL,
    LABEL_FONT,
    NEUTRAL_FILL,
    OK_FILL,
    TITLE_FONT,
    write_dataframe_sheet,
)

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.worksheet.worksheet import Worksheet

#: Nomes fixos dos artefatos (a ficha do requisito os cita por nome).
JSON_REPORT_NAME = "comparacao_report.json"
XLSX_NAME = "divergencias.xlsx"
ONLY_A_CSV_NAME = "somente_a.csv"
ONLY_B_CSV_NAME = "somente_b.csv"

#: Coluna anexada aos CSVs/abas com a linha fisica de origem.
ROW_COLUMN = "linha"

_SHEET_RESUMO = "Resumo"
_SHEET_DIVERGENCIAS = "Divergencias"
_SHEET_SOMENTE_A = "Somente em A"
_SHEET_SOMENTE_B = "Somente em B"
_SHEET_CONFLITOS = "Conflitos de chave"

_LABEL_DIFF_COLUMN = "coluna divergente"
_LABEL_VALUE_A = "valor em A"
_LABEL_VALUE_B = "valor em B"
_LABEL_ROW_A = "linha em A"
_LABEL_ROW_B = "linha em B"

#: Rotulos legiveis das categorias no resumo.
CATEGORY_LABELS: dict[str, str] = {
    "identico": "Identicos",
    "divergente": "Divergentes",
    "somente_a": "Somente em A",
    "somente_b": "Somente em B",
    "conflitos": "Conflitos de chave",
}

_CONFLICT_LABELS: dict[str, str] = {
    "chave_duplicada": "chave repetida na fonte (nao pareada)",
    "chave_vazia": "registro sem chave (nao pareado)",
}


# ============================================================
# Relatorio JSON
# ============================================================


def _key_as_dict(result: ComparisonResult, key: tuple[str, ...]) -> dict[str, str]:
    """Chave como ``{coluna: valor}`` (vazio quando o registro nao tem chave)."""
    return dict(zip(result.key_columns, key, strict=False))


def _record_payload(result: ComparisonResult, record: RecordComparison) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "chave": _key_as_dict(result, record.key),
        "linha_a": record.row_a,
        "linha_b": record.row_b,
    }
    if record.differences:
        payload["diferencas"] = [
            {"coluna": d.column, "a": d.value_a, "b": d.value_b} for d in record.differences
        ]
    return payload


def _conflict_payload(result: ComparisonResult, conflict: KeyConflict) -> dict[str, Any]:
    return {
        "chave": _key_as_dict(result, conflict.key),
        "motivo": conflict.reason,
        "descricao": _CONFLICT_LABELS.get(conflict.reason, conflict.reason),
        "linhas_a": list(conflict.rows_a),
        "linhas_b": list(conflict.rows_b),
    }


def build_report_payload(
    result: ComparisonResult,
    meta: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """
    Monta o dicionario do ``comparacao_report.json``.

    ``meta`` entra no topo (task, status, timestamps) — do mesmo jeito que o
    relatorio da auditoria faz, para que os dois artefatos sejam lidos com o
    mesmo codigo do outro lado.

    Registros identicos aparecem apenas nos contadores: listar cada um
    inflaria o arquivo sem dizer nada que o operador precise decidir.
    """
    payload: dict[str, Any] = dict(meta or {})
    payload.update(
        {
            "requisito": "RF-REC-001",
            "fonte_a": _source_payload(result, "a"),
            "fonte_b": _source_payload(result, "b"),
            "chave": list(result.key_columns),
            "colunas_comparadas": list(result.compared_columns),
            "normalizacoes": list(result.normalizations),
            "resumo": result.counts,
            "divergencias": [_record_payload(result, r) for r in result.by_category("divergente")],
            "somente_a": [_record_payload(result, r) for r in result.by_category("somente_a")],
            "somente_b": [_record_payload(result, r) for r in result.by_category("somente_b")],
            "conflitos": [_conflict_payload(result, c) for c in result.conflicts],
            "avisos": [
                {"codigo": w.code, "mensagem": w.message, "coluna": w.column}
                for w in result.warnings
            ],
            "limitacoes": list(result.limitations),
        }
    )
    return payload


def _source_payload(result: ComparisonResult, side: str) -> dict[str, Any]:
    info = result.source_a if side == "a" else result.source_b
    return {
        "arquivo": info.name,
        "aba": info.sheet,
        "linhas": info.rows,
        "primeira_linha_de_dados": info.first_data_row,
        "colunas": list(info.columns),
    }


def write_json_report(
    result: ComparisonResult,
    path: Path,
    meta: dict[str, Any] | None = None,
) -> None:
    """Grava o relatorio JSON (UTF-8, acentos preservados)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(build_report_payload(result, meta), handle, indent=2, ensure_ascii=False)


# ============================================================
# CSVs de exclusivos
# ============================================================


def _unique_column_name(frame: pd.DataFrame, base: str) -> str:
    """Nome de coluna que ainda nao existe no DataFrame (nao sobrescreve dado)."""
    nome = base
    existentes = {str(c) for c in frame.columns}
    while nome in existentes:
        nome = f"{nome}_"
    return nome


def _rows_frame(
    result: ComparisonResult,
    table: TableSource,
    category: Category,
) -> pd.DataFrame:
    """Linhas inteiras da fonte para uma categoria, + a linha fisica de origem."""
    registros = result.by_category(category)
    linhas = [r.row_a if category == "somente_a" else r.row_b for r in registros]
    posicoes = [linha - table.first_data_row for linha in linhas if linha is not None]

    subset = table.frame.iloc[posicoes].copy()
    subset[_unique_column_name(table.frame, ROW_COLUMN)] = [
        linha for linha in linhas if linha is not None
    ]
    return subset


def write_only_csvs(
    result: ComparisonResult,
    table_a: TableSource,
    table_b: TableSource,
    out_dir: Path,
) -> tuple[Path, Path]:
    """
    Grava ``somente_a.csv`` e ``somente_b.csv`` com as linhas inteiras.

    UTF-8 com BOM: e o que faz o Excel no Windows abrir os acentos certos.
    Os arquivos sao criados mesmo quando vazios (so o cabecalho) — a
    ausencia de um artefato esperado geraria duvida sobre a execucao.
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    destinos = (out_dir / ONLY_A_CSV_NAME, out_dir / ONLY_B_CSV_NAME)
    saidas: tuple[tuple[Path, TableSource, Category], ...] = (
        (destinos[0], table_a, "somente_a"),
        (destinos[1], table_b, "somente_b"),
    )
    for destino, tabela, categoria in saidas:
        _rows_frame(result, tabela, categoria).to_csv(destino, index=False, encoding="utf-8-sig")
    return destinos


# ============================================================
# Planilha de divergencias
# ============================================================


def _divergences_frame(result: ComparisonResult) -> pd.DataFrame:
    """Uma linha por CELULA divergente — o formato que se filtra no Excel."""
    linhas: list[dict[str, Any]] = []
    for record in result.by_category("divergente"):
        for diff in record.differences:
            linha: dict[str, Any] = _key_as_dict(result, record.key)
            linha[_LABEL_DIFF_COLUMN] = diff.column
            linha[_LABEL_VALUE_A] = diff.value_a
            linha[_LABEL_VALUE_B] = diff.value_b
            linha[_LABEL_ROW_A] = record.row_a
            linha[_LABEL_ROW_B] = record.row_b
            linhas.append(linha)

    colunas = [
        *result.key_columns,
        _LABEL_DIFF_COLUMN,
        _LABEL_VALUE_A,
        _LABEL_VALUE_B,
        _LABEL_ROW_A,
        _LABEL_ROW_B,
    ]
    return pd.DataFrame(linhas, columns=colunas)


def _conflicts_frame(result: ComparisonResult) -> pd.DataFrame:
    linhas = [
        {
            **_key_as_dict(result, conflict.key),
            "motivo": _CONFLICT_LABELS.get(conflict.reason, conflict.reason),
            "linhas em A": ", ".join(str(n) for n in conflict.rows_a),
            "linhas em B": ", ".join(str(n) for n in conflict.rows_b),
        }
        for conflict in result.conflicts
    ]
    colunas = [*result.key_columns, "motivo", "linhas em A", "linhas em B"]
    return pd.DataFrame(linhas, columns=colunas)


def _build_resumo(ws: Worksheet, result: ComparisonResult) -> None:
    ws["A1"] = "Comparacao de planilhas"
    ws["A1"].font = TITLE_FONT

    contagem = result.counts
    linhas: list[tuple[str, object]] = [
        ("Fonte A", f"{result.source_a.name} ({result.source_a.rows} registros)"),
        ("Fonte B", f"{result.source_b.name} ({result.source_b.rows} registros)"),
        ("Chave", ", ".join(result.key_columns)),
        ("Colunas comparadas", ", ".join(result.compared_columns) or "(nenhuma)"),
        ("Normalizacoes", ", ".join(result.normalizations) or "(nenhuma)"),
        (CATEGORY_LABELS["identico"], contagem["identico"]),
        (CATEGORY_LABELS["divergente"], contagem["divergente"]),
        (CATEGORY_LABELS["somente_a"], contagem["somente_a"]),
        (CATEGORY_LABELS["somente_b"], contagem["somente_b"]),
        (CATEGORY_LABELS["conflitos"], contagem["conflitos"]),
    ]
    destaques = {
        CATEGORY_LABELS["identico"]: OK_FILL,
        CATEGORY_LABELS["divergente"]: ERROR_FILL,
        CATEGORY_LABELS["somente_a"]: NEUTRAL_FILL,
        CATEGORY_LABELS["somente_b"]: NEUTRAL_FILL,
        CATEGORY_LABELS["conflitos"]: ERROR_FILL,
    }

    linha_atual = 3
    for rotulo, valor in linhas:
        ws.cell(row=linha_atual, column=1, value=rotulo).font = LABEL_FONT
        celula = ws.cell(row=linha_atual, column=2, value=valor)
        celula.font = DATA_FONT
        if rotulo in destaques:
            ws.cell(row=linha_atual, column=1).fill = destaques[rotulo]
            celula.fill = destaques[rotulo]
        linha_atual += 1

    linha_atual = _write_bullet_block(
        ws, linha_atual + 1, "Avisos", [w.message for w in result.warnings] or ["(nenhum)"]
    )
    _write_bullet_block(ws, linha_atual + 1, "Limitacoes desta comparacao", result.limitations)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90


def _write_bullet_block(
    ws: Worksheet,
    start_row: int,
    title: str,
    items: tuple[str, ...] | list[str],
) -> int:
    """Escreve um bloco de titulo + itens e devolve a proxima linha livre."""
    ws.cell(row=start_row, column=1, value=title).font = LABEL_FONT
    linha = start_row + 1
    for item in items:
        ws.cell(row=linha, column=2, value=f"- {item}").font = DATA_FONT
        linha += 1
    return linha


def write_divergencias_xlsx(
    result: ComparisonResult,
    table_a: TableSource,
    table_b: TableSource,
    path: Path,
) -> None:
    """Gera ``divergencias.xlsx`` com uma aba por categoria."""
    wb = Workbook()
    resumo = wb.active
    if resumo is None:  # pragma: no cover - openpyxl sempre cria a primeira aba
        resumo = wb.create_sheet()
    resumo.title = _SHEET_RESUMO
    _build_resumo(resumo, result)

    write_dataframe_sheet(wb.create_sheet(_SHEET_DIVERGENCIAS), _divergences_frame(result))
    write_dataframe_sheet(
        wb.create_sheet(_SHEET_SOMENTE_A), _rows_frame(result, table_a, "somente_a")
    )
    write_dataframe_sheet(
        wb.create_sheet(_SHEET_SOMENTE_B), _rows_frame(result, table_b, "somente_b")
    )
    write_dataframe_sheet(wb.create_sheet(_SHEET_CONFLITOS), _conflicts_frame(result))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ============================================================
# Resumo textual (terminal)
# ============================================================


def format_key(result: ComparisonResult, key: tuple[str, ...]) -> str:
    """Chave legivel em uma linha: ``codigo=A-1; loja=SP``."""
    if not key:
        return "(sem chave)"
    return "; ".join(
        f"{coluna}={valor}" for coluna, valor in zip(result.key_columns, key, strict=False)
    )


def _record_line(result: ComparisonResult, record: RecordComparison) -> list[str]:
    onde = f"linha {record.row_a} em A, linha {record.row_b} em B"
    linhas = [f"  - {format_key(result, record.key)} ({onde})"]
    linhas.extend(f"      {d.column}: '{d.value_a}'  ->  '{d.value_b}'" for d in record.differences)
    return linhas


def generate_summary(result: ComparisonResult, *, max_rows: int = 10) -> str:
    """
    Resumo da comparacao para o terminal.

    Mostra os contadores, as primeiras divergencias (com os valores dos dois
    lados), os conflitos e, sempre, os avisos e as limitacoes — quem le
    precisa saber tambem o que a comparacao NAO olhou.
    """
    contagem = result.counts
    partes: list[str] = [
        f"Comparacao por chave: {', '.join(result.key_columns)}",
        f"  Fonte A: {result.source_a.name} ({result.source_a.rows} registros)",
        f"  Fonte B: {result.source_b.name} ({result.source_b.rows} registros)",
        f"  Colunas comparadas: {', '.join(result.compared_columns) or '(nenhuma)'}",
        "",
    ]
    partes.extend(
        f"  {CATEGORY_LABELS[categoria]:<20} {contagem[categoria]}"
        for categoria in ("identico", "divergente", "somente_a", "somente_b", "conflitos")
    )

    divergentes = result.by_category("divergente")
    if divergentes:
        partes.append("")
        partes.append(f"Divergencias ({len(divergentes)}):")
        for record in divergentes[:max_rows]:
            partes.extend(_record_line(result, record))
        if len(divergentes) > max_rows:
            partes.append(f"  ... e mais {len(divergentes) - max_rows}")

    exclusivos: tuple[tuple[Category, str], ...] = (
        ("somente_a", "Somente em A"),
        ("somente_b", "Somente em B"),
    )
    for categoria, rotulo in exclusivos:
        registros = result.by_category(categoria)
        if not registros:
            continue
        partes.append("")
        partes.append(f"{rotulo} ({len(registros)}):")
        for record in registros[:max_rows]:
            linha = record.row_a if categoria == "somente_a" else record.row_b
            partes.append(f"  - {format_key(result, record.key)} (linha {linha})")
        if len(registros) > max_rows:
            partes.append(f"  ... e mais {len(registros) - max_rows}")

    if result.conflicts:
        partes.append("")
        partes.append(f"Conflitos de chave ({len(result.conflicts)}) — nao pareados:")
        for conflito in result.conflicts[:max_rows]:
            linhas_a = ", ".join(str(n) for n in conflito.rows_a) or "-"
            linhas_b = ", ".join(str(n) for n in conflito.rows_b) or "-"
            partes.append(
                f"  - {format_key(result, conflito.key)}: "
                f"{_CONFLICT_LABELS.get(conflito.reason, conflito.reason)} "
                f"(A: {linhas_a} | B: {linhas_b})"
            )

    if result.warnings:
        partes.append("")
        partes.append("Avisos:")
        partes.extend(f"  - {w.message}" for w in result.warnings)

    if result.limitations:
        partes.append("")
        partes.append("Limitacoes desta comparacao:")
        partes.extend(f"  - {limite}" for limite in result.limitations)

    return "\n".join(partes)


# ============================================================
# Entry point
# ============================================================


def write_comparison_artifacts(
    result: ComparisonResult,
    table_a: TableSource,
    table_b: TableSource,
    out_dir: Path,
    meta: dict[str, Any] | None = None,
) -> tuple[Path, ...]:
    """
    Gera os quatro artefatos em ``out_dir`` (nomes fixos).

    Returns:
        Os caminhos criados, na ordem: JSON, XLSX, somente_a, somente_b.
    """
    out_dir.mkdir(parents=True, exist_ok=True)

    json_path = out_dir / JSON_REPORT_NAME
    write_json_report(result, json_path, meta)

    xlsx_path = out_dir / XLSX_NAME
    write_divergencias_xlsx(result, table_a, table_b, xlsx_path)

    csv_a, csv_b = write_only_csvs(result, table_a, table_b, out_dir)
    return (json_path, xlsx_path, csv_a, csv_b)


__all__ = [
    "CATEGORY_LABELS",
    "JSON_REPORT_NAME",
    "ONLY_A_CSV_NAME",
    "ONLY_B_CSV_NAME",
    "ROW_COLUMN",
    "XLSX_NAME",
    "build_report_payload",
    "format_key",
    "generate_summary",
    "write_comparison_artifacts",
    "write_divergencias_xlsx",
    "write_json_report",
    "write_only_csvs",
]
