"""
Artefatos da transferencia entre planilhas (RF-REC-003).

- ``planilha_enriquecida.xlsx`` — o destino com os campos autorizados
  preenchidos, mais as abas de evidencia (Transferencias, Nao encontrados,
  Resumo).
- ``planilha_enriquecida.csv``  — o mesmo conteudo para quem segue no fluxo.
- ``nao_encontrados.csv``       — as chaves buscadas que a fonte nao tinha.
- ``transferencia_report.json`` — tudo em formato de maquina.

O destino original nao e tocado em momento algum.
"""

from __future__ import annotations

import json
from typing import TYPE_CHECKING, Any

import pandas as pd
from openpyxl import Workbook

from autotarefas.reconcile.artifacts import format_key
from autotarefas.reconcile.transfer import TransferResult, source_field_warnings
from autotarefas.tasks.xlsx_style import (
    DATA_FONT,
    ERROR_FILL,
    LABEL_FONT,
    NEUTRAL_FILL,
    OK_FILL,
    TITLE_FONT,
    write_dataframe_sheet,
)

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.worksheet.worksheet import Worksheet

    from autotarefas.reconcile.result import TableSource

ENRICHED_XLSX_NAME = "planilha_enriquecida.xlsx"
ENRICHED_CSV_NAME = "planilha_enriquecida.csv"
NOT_FOUND_CSV_NAME = "nao_encontrados.csv"
JSON_REPORT_NAME = "transferencia_report.json"

_SHEET_BASE = "Planilha enriquecida"
_SHEET_TRANSFERENCIAS = "Transferencias"
_SHEET_NAO_ENCONTRADOS = "Nao encontrados"
_SHEET_RESUMO = "Resumo"

ACTION_LABELS: dict[str, str] = {
    "preenchido": "campo vazio preenchido pela fonte",
    "atualizado": "valor substituido pelo da fonte",
    "mantido": "valor do destino preservado",
    "sem_valor_na_fonte": "a fonte nao tinha valor (destino intacto)",
}


def enriched_frame(result: TransferResult) -> pd.DataFrame:
    """O destino ja enriquecido, na ordem original das colunas."""
    return pd.DataFrame(list(result.rows), columns=list(result.columns), dtype=str).fillna("")


def transfers_frame(result: TransferResult) -> pd.DataFrame:
    """Uma linha por celula tocada (ou deliberadamente nao tocada)."""
    chaves = result.comparison.key_columns
    linhas: list[dict[str, Any]] = [
        {
            **dict(zip(chaves, transfer.key, strict=False)),
            "linha no destino": transfer.row,
            "coluna": transfer.column,
            "antes": transfer.before,
            "depois": transfer.after,
            "acao": ACTION_LABELS.get(transfer.action, transfer.action),
        }
        for transfer in result.transfers
    ]
    colunas = [*chaves, "linha no destino", "coluna", "antes", "depois", "acao"]
    return pd.DataFrame(linhas, columns=colunas)


def not_found_frame(result: TransferResult) -> pd.DataFrame:
    """As chaves buscadas sem correspondente na fonte."""
    chaves = result.comparison.key_columns
    linhas: list[dict[str, Any]] = [
        {
            **dict(zip(chaves, item.key, strict=False)),
            "linha no destino": item.row,
            "situacao": "sem correspondente na fonte",
        }
        for item in result.not_found
    ]
    return pd.DataFrame(linhas, columns=[*chaves, "linha no destino", "situacao"])


def build_report_payload(
    result: TransferResult,
    source: TableSource | None = None,
    meta: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Monta o dicionario do ``transferencia_report.json``."""
    comparacao = result.comparison
    payload: dict[str, Any] = dict(meta or {})
    payload.update(
        {
            "requisito": "RF-REC-003",
            "destino": {"arquivo": comparacao.source_a.name, "linhas": comparacao.source_a.rows},
            "fonte": {"arquivo": comparacao.source_b.name, "linhas": comparacao.source_b.rows},
            "chave": list(comparacao.key_columns),
            "politica": {
                "campos_autorizados": list(result.policy.fields),
                "somente_campos_vazios": result.policy.only_empty,
                "marca_origem": result.policy.mark_origin,
                "descricao": result.policy.describe(),
            },
            "resumo": result.counts,
            "transferencias": [
                {
                    "chave": dict(zip(comparacao.key_columns, t.key, strict=False)),
                    "linha": t.row,
                    "coluna": t.column,
                    "antes": t.before,
                    "depois": t.after,
                    "acao": t.action,
                }
                for t in result.transfers
            ],
            "nao_encontrados": [
                {
                    "chave": dict(zip(comparacao.key_columns, item.key, strict=False)),
                    "linha": item.row,
                }
                for item in result.not_found
            ],
            "so_na_fonte": [
                dict(zip(comparacao.key_columns, chave, strict=False))
                for chave in result.unmatched_source
            ],
            "avisos": [
                *(source_field_warnings(source, result.policy) if source is not None else []),
                *(w.message for w in comparacao.warnings),
            ],
            "limitacoes": list(result.limitations),
        }
    )
    return payload


def write_json_report(
    result: TransferResult,
    path: Path,
    source: TableSource | None = None,
    meta: dict[str, Any] | None = None,
) -> None:
    """Grava o relatorio da transferencia em JSON."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(build_report_payload(result, source, meta), handle, indent=2, ensure_ascii=False)


def _build_resumo(ws: Worksheet, result: TransferResult) -> None:
    ws["A1"] = "Transferencia entre planilhas"
    ws["A1"].font = TITLE_FONT

    contagem = result.counts
    linhas: list[tuple[str, object]] = [
        ("Destino", result.comparison.source_a.name),
        ("Fonte", result.comparison.source_b.name),
        ("Chave", ", ".join(result.comparison.key_columns)),
        ("Politica", result.policy.describe()),
        ("Registros no destino", contagem["registros_no_destino"]),
        ("Campos preenchidos", contagem["preenchidos"]),
        ("Campos atualizados", contagem["atualizados"]),
        ("Valores preservados", contagem["mantidos"]),
        ("Sem valor na fonte", contagem["sem_valor_na_fonte"]),
        ("Nao encontrados", contagem["nao_encontrados"]),
        ("Registros so na fonte", contagem["so_na_fonte"]),
    ]
    destaques = {
        "Campos preenchidos": OK_FILL,
        "Campos atualizados": NEUTRAL_FILL,
        "Nao encontrados": ERROR_FILL,
    }

    linha = 3
    for rotulo, valor in linhas:
        ws.cell(row=linha, column=1, value=rotulo).font = LABEL_FONT
        celula = ws.cell(row=linha, column=2, value=valor)
        celula.font = DATA_FONT
        if rotulo in destaques:
            ws.cell(row=linha, column=1).fill = destaques[rotulo]
            celula.fill = destaques[rotulo]
        linha += 1

    ws.cell(row=linha + 1, column=1, value="Limitacoes desta transferencia").font = LABEL_FONT
    linha += 2
    for limite in result.limitations:
        ws.cell(row=linha, column=2, value=f"- {limite}").font = DATA_FONT
        linha += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90


def write_enriched_xlsx(result: TransferResult, path: Path) -> None:
    """Gera ``planilha_enriquecida.xlsx`` com as abas de evidencia."""
    wb = Workbook()
    base = wb.active
    if base is None:  # pragma: no cover
        base = wb.create_sheet()
    base.title = _SHEET_BASE
    write_dataframe_sheet(base, enriched_frame(result))
    write_dataframe_sheet(wb.create_sheet(_SHEET_TRANSFERENCIAS), transfers_frame(result))
    write_dataframe_sheet(wb.create_sheet(_SHEET_NAO_ENCONTRADOS), not_found_frame(result))
    _build_resumo(wb.create_sheet(_SHEET_RESUMO), result)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def generate_summary(result: TransferResult, *, max_rows: int = 10) -> str:
    """Resumo da transferencia para o terminal."""
    contagem = result.counts
    comparacao = result.comparison
    partes: list[str] = [
        f"Transferencia por chave: {', '.join(comparacao.key_columns)}",
        f"  Destino: {comparacao.source_a.name} ({comparacao.source_a.rows} registros)",
        f"  Fonte:   {comparacao.source_b.name} ({comparacao.source_b.rows} registros)",
        f"  {result.policy.describe()}",
        "",
        f"  Campos preenchidos    {contagem['preenchidos']}",
        f"  Campos atualizados    {contagem['atualizados']}",
        f"  Valores preservados   {contagem['mantidos']}",
        f"  Sem valor na fonte    {contagem['sem_valor_na_fonte']}",
        f"  Nao encontrados       {contagem['nao_encontrados']}",
    ]

    alteracoes = [t for t in result.transfers if t.changed]
    if alteracoes:
        partes.append("")
        partes.append(f"Alteracoes ({len(alteracoes)}):")
        for transfer in alteracoes[:max_rows]:
            partes.append(
                f"  - {format_key(comparacao, transfer.key)} | {transfer.column}: "
                f"'{transfer.before}'  ->  '{transfer.after}'"
            )
        if len(alteracoes) > max_rows:
            partes.append(f"  ... e mais {len(alteracoes) - max_rows}")

    if result.not_found:
        partes.append("")
        partes.append(f"Sem correspondente na fonte ({len(result.not_found)}):")
        for item in result.not_found[:max_rows]:
            partes.append(f"  - {format_key(comparacao, item.key)} (linha {item.row})")
        if len(result.not_found) > max_rows:
            partes.append(f"  ... e mais {len(result.not_found) - max_rows}")

    partes.append("")
    partes.append("Limitacoes desta transferencia:")
    partes.extend(f"  - {limite}" for limite in result.limitations)
    return "\n".join(partes)


def write_transfer_artifacts(
    result: TransferResult,
    out_dir: Path,
    source: TableSource | None = None,
    meta: dict[str, Any] | None = None,
) -> tuple[Path, ...]:
    """
    Gera os artefatos da transferencia em ``out_dir`` (nomes fixos).

    Returns:
        Caminhos criados: XLSX enriquecido, CSV enriquecido, nao encontrados,
        relatorio JSON.
    """
    out_dir.mkdir(parents=True, exist_ok=True)

    xlsx = out_dir / ENRICHED_XLSX_NAME
    write_enriched_xlsx(result, xlsx)

    csv = out_dir / ENRICHED_CSV_NAME
    enriched_frame(result).to_csv(csv, index=False, encoding="utf-8-sig")

    nao_encontrados = out_dir / NOT_FOUND_CSV_NAME
    not_found_frame(result).to_csv(nao_encontrados, index=False, encoding="utf-8-sig")

    relatorio = out_dir / JSON_REPORT_NAME
    write_json_report(result, relatorio, source, meta)

    return (xlsx, csv, nao_encontrados, relatorio)


__all__ = [
    "ACTION_LABELS",
    "ENRICHED_CSV_NAME",
    "ENRICHED_XLSX_NAME",
    "JSON_REPORT_NAME",
    "NOT_FOUND_CSV_NAME",
    "build_report_payload",
    "enriched_frame",
    "generate_summary",
    "not_found_frame",
    "transfers_frame",
    "write_enriched_xlsx",
    "write_json_report",
    "write_transfer_artifacts",
]
