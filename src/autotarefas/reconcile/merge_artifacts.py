"""
Artefatos da reconciliacao (RF-REC-002).

Tres saidas, tres publicos:

- ``base_conciliada.xlsx``        — a base resultante, pronta para uso, com
  uma aba de decisoes ao lado (valor escolhido, fonte e regra por celula).
- ``conflitos_para_revisao.xlsx`` — o que o AutoTarefas se recusou a decidir:
  divergencia sem autorizacao, chave duplicada, chave vazia.
- ``reconciliacao_report.json``   — o mesmo conteudo em formato de maquina,
  com a politica declarada e as limitacoes.

A base conciliada e sempre um ARQUIVO NOVO: nenhuma das fontes e tocada.
"""

from __future__ import annotations

import json
from typing import TYPE_CHECKING, Any

import pandas as pd
from openpyxl import Workbook

from autotarefas.reconcile.artifacts import format_key
from autotarefas.reconcile.merge import ReconciliationResult
from autotarefas.tasks.xlsx_style import (
    DATA_FONT,
    ERROR_FILL,
    LABEL_FONT,
    NEUTRAL_FILL,
    OK_FILL,
    TITLE_FONT,
    write_dataframe_sheet,
)

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.worksheet.worksheet import Worksheet

#: Nomes fixos dos artefatos da reconciliacao.
BASE_XLSX_NAME = "base_conciliada.xlsx"
REVIEW_XLSX_NAME = "conflitos_para_revisao.xlsx"
JSON_REPORT_NAME = "reconciliacao_report.json"
BASE_CSV_NAME = "base_conciliada.csv"

_SHEET_BASE = "Base conciliada"
_SHEET_DECISOES = "Decisoes"
_SHEET_RESUMO = "Resumo"
_SHEET_REVISAO = "Para revisao"

#: Rotulos legiveis das regras e dos motivos (o operador nao le enum).
RULE_LABELS: dict[str, str] = {
    "fonte_principal": "valor da fonte principal (sem autorizacao para mudar)",
    "campo_autorizado": "atualizado pela complementar (campo autorizado)",
    "tolerancia": "diferenca dentro da tolerancia declarada",
    "registro_exclusivo": "registro existia so na fonte complementar",
}

REVIEW_LABELS: dict[str, str] = {
    "divergencia_nao_autorizada": "divergencia em campo nao autorizado",
    "chave_duplicada": "chave repetida na fonte (nao pareada)",
    "chave_vazia": "registro sem chave (nao pareado)",
    "registro_novo_nao_incluido": "registro novo da complementar (nao incluido por opcao)",
}


# ============================================================
# Estruturas tabulares
# ============================================================


def base_frame(result: ReconciliationResult) -> pd.DataFrame:
    """A base conciliada como DataFrame (colunas da fonte principal + origem)."""
    linhas = [dict(record.values) for record in result.records]
    frame = pd.DataFrame(linhas, columns=list(result.columns), dtype=str)
    frame = frame.fillna("")
    frame["_origem"] = [record.origin for record in result.records]
    return frame


def decisions_frame(result: ReconciliationResult) -> pd.DataFrame:
    """Uma linha por decisao automatica: chave, coluna, valor, fonte e regra."""
    linhas: list[dict[str, Any]] = [
        {
            **dict(zip(result.comparison.key_columns, record.key, strict=False)),
            "coluna": decision.column,
            "valor escolhido": decision.value,
            "valor descartado": decision.other_value,
            "fonte": decision.source.upper(),
            "regra": RULE_LABELS.get(decision.rule, decision.rule),
        }
        for record in result.records
        for decision in record.decisions
    ]
    colunas = [
        *result.comparison.key_columns,
        "coluna",
        "valor escolhido",
        "valor descartado",
        "fonte",
        "regra",
    ]
    return pd.DataFrame(linhas, columns=colunas)


def review_frame(result: ReconciliationResult) -> pd.DataFrame:
    """Os casos que ficaram para uma pessoa decidir."""
    linhas: list[dict[str, Any]] = [
        {
            **dict(zip(result.comparison.key_columns, item.key, strict=False)),
            "motivo": REVIEW_LABELS.get(item.reason, item.reason),
            "coluna": item.column or "",
            "valor em A": item.value_a,
            "valor em B": item.value_b,
            "linhas em A": ", ".join(str(n) for n in item.rows_a),
            "linhas em B": ", ".join(str(n) for n in item.rows_b),
        }
        for item in result.review
    ]
    colunas = [
        *result.comparison.key_columns,
        "motivo",
        "coluna",
        "valor em A",
        "valor em B",
        "linhas em A",
        "linhas em B",
    ]
    return pd.DataFrame(linhas, columns=colunas)


# ============================================================
# Relatorio JSON
# ============================================================


def build_report_payload(
    result: ReconciliationResult,
    meta: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Monta o dicionario do ``reconciliacao_report.json``."""
    comparacao = result.comparison
    payload: dict[str, Any] = dict(meta or {})
    payload.update(
        {
            "requisito": "RF-REC-002",
            "fonte_a": {"arquivo": comparacao.source_a.name, "linhas": comparacao.source_a.rows},
            "fonte_b": {"arquivo": comparacao.source_b.name, "linhas": comparacao.source_b.rows},
            "politica": {
                "fonte_principal": result.policy.primary,
                "campos_atualizaveis": list(result.policy.updatable_columns),
                "incluir_novos_da_complementar": result.policy.include_new,
                "descricao": result.policy.describe(),
            },
            "chave": list(comparacao.key_columns),
            "colunas_da_base": list(result.columns),
            "tolerancias": list(comparacao.tolerances),
            "resumo": {**comparacao.counts, **result.counts},
            "decisoes": [
                {
                    "chave": dict(zip(comparacao.key_columns, record.key, strict=False)),
                    "coluna": decision.column,
                    "valor": decision.value,
                    "valor_descartado": decision.other_value,
                    "fonte": decision.source,
                    "regra": decision.rule,
                }
                for record in result.records
                for decision in record.decisions
            ],
            "para_revisao": [
                {
                    "chave": dict(zip(comparacao.key_columns, item.key, strict=False)),
                    "motivo": item.reason,
                    "coluna": item.column,
                    "valor_a": item.value_a,
                    "valor_b": item.value_b,
                    "linhas_a": list(item.rows_a),
                    "linhas_b": list(item.rows_b),
                }
                for item in result.review
            ],
            "limitacoes": [*result.limitations, *comparacao.limitations],
        }
    )
    return payload


def write_json_report(
    result: ReconciliationResult,
    path: Path,
    meta: dict[str, Any] | None = None,
) -> None:
    """Grava o relatorio de decisoes em JSON."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(build_report_payload(result, meta), handle, indent=2, ensure_ascii=False)


# ============================================================
# Planilhas
# ============================================================


def _build_resumo(ws: Worksheet, result: ReconciliationResult) -> None:
    ws["A1"] = "Reconciliacao de bases"
    ws["A1"].font = TITLE_FONT

    comparacao = result.comparison
    contagem = {**comparacao.counts, **result.counts}
    linhas: list[tuple[str, object]] = [
        ("Fonte A", comparacao.source_a.name),
        ("Fonte B", comparacao.source_b.name),
        ("Chave", ", ".join(comparacao.key_columns)),
        ("Politica", result.policy.describe()),
        ("Tolerancias", "; ".join(comparacao.tolerances) or "(nenhuma)"),
        ("Registros conciliados", contagem["conciliados"]),
        ("Registros atualizados", contagem["registros_atualizados"]),
        ("Campos atualizados", contagem["campos_atualizados"]),
        ("Diferencas toleradas", contagem["toleradas"]),
        ("Casos para revisao", contagem["para_revisao"]),
    ]
    destaques = {
        "Registros conciliados": OK_FILL,
        "Campos atualizados": NEUTRAL_FILL,
        "Casos para revisao": ERROR_FILL,
    }

    linha = 3
    for rotulo, valor in linhas:
        ws.cell(row=linha, column=1, value=rotulo).font = LABEL_FONT
        celula = ws.cell(row=linha, column=2, value=valor)
        celula.font = DATA_FONT
        if rotulo in destaques:
            ws.cell(row=linha, column=1).fill = destaques[rotulo]
            celula.fill = destaques[rotulo]
        linha += 1

    ws.cell(row=linha + 1, column=1, value="Limitacoes desta reconciliacao").font = LABEL_FONT
    linha += 2
    for limite in result.limitations:
        ws.cell(row=linha, column=2, value=f"- {limite}").font = DATA_FONT
        linha += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90


def write_base_xlsx(result: ReconciliationResult, path: Path) -> None:
    """Gera ``base_conciliada.xlsx`` (base + decisoes + resumo)."""
    wb = Workbook()
    base = wb.active
    if base is None:  # pragma: no cover - openpyxl sempre cria a primeira aba
        base = wb.create_sheet()
    base.title = _SHEET_BASE
    write_dataframe_sheet(base, base_frame(result))
    write_dataframe_sheet(wb.create_sheet(_SHEET_DECISOES), decisions_frame(result))
    _build_resumo(wb.create_sheet(_SHEET_RESUMO), result)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_review_xlsx(result: ReconciliationResult, path: Path) -> None:
    """Gera ``conflitos_para_revisao.xlsx``."""
    wb = Workbook()
    ws = wb.active
    if ws is None:  # pragma: no cover
        ws = wb.create_sheet()
    ws.title = _SHEET_REVISAO
    write_dataframe_sheet(ws, review_frame(result))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ============================================================
# Resumo textual e entry point
# ============================================================


def generate_summary(result: ReconciliationResult, *, max_rows: int = 10) -> str:
    """Resumo da reconciliacao para o terminal."""
    comparacao = result.comparison
    contagem = {**comparacao.counts, **result.counts}
    partes: list[str] = [
        f"Reconciliacao por chave: {', '.join(comparacao.key_columns)}",
        f"  {result.policy.describe()}",
        f"  Tolerancias: {'; '.join(comparacao.tolerances) or '(nenhuma)'}",
        "",
        f"  Registros conciliados   {contagem['conciliados']}",
        f"  Registros atualizados   {contagem['registros_atualizados']}",
        f"  Campos atualizados      {contagem['campos_atualizados']}",
        f"  Diferencas toleradas    {contagem['toleradas']}",
        f"  Casos para revisao      {contagem['para_revisao']}",
    ]

    atualizacoes = [
        (record, decision)
        for record in result.records
        for decision in record.decisions
        if decision.rule == "campo_autorizado"
    ]
    if atualizacoes:
        partes.append("")
        partes.append(f"Atualizacoes aplicadas ({len(atualizacoes)}):")
        for record, decision in atualizacoes[:max_rows]:
            partes.append(
                f"  - {format_key(comparacao, record.key)} | {decision.column}: "
                f"'{decision.other_value}'  ->  '{decision.value}' "
                f"(fonte {decision.source.upper()})"
            )
        if len(atualizacoes) > max_rows:
            partes.append(f"  ... e mais {len(atualizacoes) - max_rows}")

    if result.review:
        partes.append("")
        partes.append(f"Para revisao ({len(result.review)}):")
        for item in result.review[:max_rows]:
            coluna = f" | {item.column}" if item.column else ""
            partes.append(
                f"  - {format_key(comparacao, item.key)}{coluna}: "
                f"{REVIEW_LABELS.get(item.reason, item.reason)}"
            )
        if len(result.review) > max_rows:
            partes.append(f"  ... e mais {len(result.review) - max_rows}")

    partes.append("")
    partes.append("Limitacoes desta reconciliacao:")
    partes.extend(f"  - {limite}" for limite in result.limitations)
    return "\n".join(partes)


def write_reconciliation_artifacts(
    result: ReconciliationResult,
    out_dir: Path,
    meta: dict[str, Any] | None = None,
) -> tuple[Path, ...]:
    """
    Gera os artefatos da reconciliacao em ``out_dir`` (nomes fixos).

    Returns:
        Caminhos criados: base XLSX, base CSV, revisao XLSX, relatorio JSON.
    """
    out_dir.mkdir(parents=True, exist_ok=True)

    base_xlsx = out_dir / BASE_XLSX_NAME
    write_base_xlsx(result, base_xlsx)

    base_csv = out_dir / BASE_CSV_NAME
    base_frame(result).to_csv(base_csv, index=False, encoding="utf-8-sig")

    revisao = out_dir / REVIEW_XLSX_NAME
    write_review_xlsx(result, revisao)

    relatorio = out_dir / JSON_REPORT_NAME
    write_json_report(result, relatorio, meta)

    return (base_xlsx, base_csv, revisao, relatorio)


__all__ = [
    "BASE_CSV_NAME",
    "BASE_XLSX_NAME",
    "JSON_REPORT_NAME",
    "REVIEW_LABELS",
    "REVIEW_XLSX_NAME",
    "RULE_LABELS",
    "base_frame",
    "build_report_payload",
    "decisions_frame",
    "generate_summary",
    "review_frame",
    "write_base_xlsx",
    "write_json_report",
    "write_reconciliation_artifacts",
    "write_review_xlsx",
]
