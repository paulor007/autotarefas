"""
`relatorio_analise.xlsx` — o laudo da execução, em abas.

Este arquivo NÃO é a planilha do cliente: é o que o AutoTarefas encontrou,
o que fez e o que deixou para uma pessoa decidir. Ele existe para que a
conferência não dependa de acreditar na tela.

    Resumo                 o que entrou, o veredito da apresentação, contadores
    Abas do arquivo        cada aba e a natureza que o diagnóstico atribuiu
    Problemas encontrados  uma linha por problema, com linha física e categoria
    Linhas para revisão    o que exige decisão humana, com o contexto
    Alterações realizadas  cada mudança aplicada (valor e apresentação)
    Antes e depois         valor anterior x posterior, célula a célula
    Indicadores confirmados  só quando a pessoa confirmou o papel das colunas

Abas sem conteúdo não são escritas: uma aba vazia sugere que algo falhou.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import TYPE_CHECKING, Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from autotarefas.tasks.xlsx_style import (
    DATA_FONT,
    ERROR_FILL,
    LABEL_FONT,
    NEUTRAL_FILL,
    OK_FILL,
    TITLE_FONT,
    write_dataframe_sheet,
)

if TYPE_CHECKING:
    from collections.abc import Mapping, Sequence
    from pathlib import Path

    from openpyxl.worksheet.worksheet import Worksheet

    from autotarefas.organize.insights import Indicator
    from autotarefas.organize.organizer import OrganizeResult
    from autotarefas.organize.presentation import PresentationAudit
    from autotarefas.organize.sheets import SheetInfo
    from autotarefas.tasks.execution_package import ReviewRow

#: Nome do artefato.
ANALYSIS_REPORT_NAME = "relatorio_analise.xlsx"

_ABA_RESUMO = "Resumo"
_ABA_ABAS = "Abas do arquivo"
_ABA_PROBLEMAS = "Problemas encontrados"
_ABA_REVISAO = "Linhas para revisao"
_ABA_ALTERACOES = "Alteracoes realizadas"
_ABA_ANTES_DEPOIS = "Antes e depois"
_ABA_INDICADORES = "Indicadores confirmados"

#: Quantas linhas do indicador viram gráfico (além disso vira poluição).
_MAX_BARRAS = 12

_VEREDITOS = {
    "organizada": "já organizada — nada a melhorar com segurança",
    "melhoravel": "pode ser melhorada",
    "ambigua": "estrutura ambígua — precisa de decisão",
}


@dataclass(frozen=True, slots=True)
class ReportInput:
    """Tudo o que o relatório precisa saber sobre a execução."""

    source_name: str
    sheet: str = ""
    rows: int = 0
    columns: int = 0
    audit: PresentationAudit | None = None
    sheets: tuple[SheetInfo, ...] = ()
    issues: Sequence[Mapping[str, Any]] = ()
    cleaning_changes: Sequence[Mapping[str, Any]] = ()
    review_rows: Sequence[ReviewRow] = ()
    organize: OrganizeResult | None = None
    indicators: Sequence[Indicator] = ()
    counters: Mapping[str, Any] = field(default_factory=dict)
    notes: Sequence[str] = ()


# ============================================================
# Abas
# ============================================================


def _resumo(ws: Worksheet, dados: ReportInput) -> None:
    ws["A1"] = "Análise e organização de planilhas"
    ws["A1"].font = TITLE_FONT

    audit = dados.audit
    organizacao = dados.organize
    linhas: list[tuple[str, object]] = [
        ("Arquivo", dados.source_name),
        ("Aba analisada", dados.sheet or "—"),
        ("Linhas de dados", dados.rows),
        ("Colunas", dados.columns),
        (
            "Apresentação",
            _VEREDITOS.get(audit.verdict, audit.verdict) if audit else "não avaliada",
        ),
        ("Problemas encontrados", len(dados.issues)),
        ("Linhas para revisão", len(dados.review_rows)),
        ("Valores corrigidos", len(dados.cleaning_changes)),
        (
            "Mudanças de apresentação",
            len(organizacao.changes) if organizacao else 0,
        ),
        ("Ordenação aplicada", (organizacao.sorted_by if organizacao else "") or "nenhuma"),
    ]
    destaques = {
        "Problemas encontrados": ERROR_FILL,
        "Linhas para revisão": ERROR_FILL,
        "Valores corrigidos": NEUTRAL_FILL,
        "Apresentação": OK_FILL,
    }

    linha = 3
    for rotulo, valor in linhas:
        ws.cell(row=linha, column=1, value=rotulo).font = LABEL_FONT
        celula = ws.cell(row=linha, column=2, value=valor)
        celula.font = DATA_FONT
        if rotulo in destaques:
            ws.cell(row=linha, column=1).fill = destaques[rotulo]
            celula.fill = destaques[rotulo]
        linha += 1

    if audit is not None:
        linha += 1
        ws.cell(row=linha, column=1, value="Critérios de apresentação").font = LABEL_FONT
        linha += 1
        for criterio in audit.criteria:
            if not criterio.applicable:
                continue
            marca = "OK" if criterio.passed else "melhorar"
            ws.cell(row=linha, column=1, value=criterio.title).font = DATA_FONT
            ws.cell(row=linha, column=2, value=f"{marca} — {criterio.detail}").font = DATA_FONT
            linha += 1

    observacoes = [
        *(dados.notes or ()),
        *(organizacao.refusals if organizacao else ()),
        *(organizacao.not_preserved if organizacao else ()),
    ]
    if observacoes:
        linha += 1
        ws.cell(row=linha, column=1, value="Observações e limites").font = LABEL_FONT
        linha += 1
        for texto in observacoes:
            ws.cell(row=linha, column=2, value=f"- {texto}").font = DATA_FONT
            linha += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95


def _abas_frame(dados: ReportInput) -> pd.DataFrame:
    linhas = [
        {
            "aba": info.name,
            "natureza": info.kind,
            "linhas": info.rows,
            "colunas": info.columns,
            "analisada": "sim" if info.name == dados.sheet else "não",
            "observação": info.reason,
        }
        for info in dados.sheets
    ]
    return pd.DataFrame(
        linhas, columns=["aba", "natureza", "linhas", "colunas", "analisada", "observação"]
    )


def _problemas_frame(dados: ReportInput) -> pd.DataFrame:
    linhas = [
        {
            "linha": issue.get("line", ""),
            "coluna": issue.get("column") or "",
            "severidade": "erro" if issue.get("severity") == "error" else "aviso",
            "problema": issue.get("message", ""),
            "valor": issue.get("value") or "",
        }
        for issue in dados.issues
    ]
    return pd.DataFrame(linhas, columns=["linha", "coluna", "severidade", "problema", "valor"])


def _revisao_frame(dados: ReportInput) -> pd.DataFrame:
    linhas = [
        {
            "linha": item.line,
            "severidade": item.severity,
            "categoria": ", ".join(item.categories),
            "grupo de duplicidade": item.duplicate_group,
            "linha canônica": item.canonical_line or "",
            "linhas relacionadas": " ".join(str(n) for n in item.related_lines),
            "motivo": " | ".join(item.reasons),
        }
        for item in dados.review_rows
    ]
    return pd.DataFrame(
        linhas,
        columns=[
            "linha",
            "severidade",
            "categoria",
            "grupo de duplicidade",
            "linha canônica",
            "linhas relacionadas",
            "motivo",
        ],
    )


def _alteracoes_frame(dados: ReportInput) -> pd.DataFrame:
    linhas: list[dict[str, Any]] = [
        {
            "natureza": "valor",
            "tipo": "correção segura",
            "onde": f"linha {mudanca.get('line')}, {mudanca.get('column')}",
            "detalhe": f"'{mudanca.get('before')}' → '{mudanca.get('after')}'",
            "motivo": ", ".join(mudanca.get("rules") or []),
        }
        for mudanca in dados.cleaning_changes
    ]
    if dados.organize is not None:
        linhas.extend(
            {
                "natureza": "apresentação",
                "tipo": mudanca.kind,
                "onde": mudanca.scope,
                "detalhe": mudanca.detail,
                "motivo": "organização confirmada",
            }
            for mudanca in dados.organize.changes
        )
        if dados.organize.sorted_by:
            linhas.append(
                {
                    "natureza": "ordem",
                    "tipo": "ordenacao_aplicada",
                    "onde": "planilha",
                    "detalhe": dados.organize.sorted_by,
                    "motivo": "ordenação confirmada",
                }
            )
    return pd.DataFrame(linhas, columns=["natureza", "tipo", "onde", "detalhe", "motivo"])


def _antes_depois_frame(dados: ReportInput) -> pd.DataFrame:
    linhas = [
        {
            "linha": mudanca.get("line"),
            "coluna": mudanca.get("column"),
            "valor anterior": mudanca.get("before"),
            "valor posterior": mudanca.get("after"),
            "tipo da alteração": "correção segura",
            "motivo": ", ".join(mudanca.get("rules") or []),
        }
        for mudanca in dados.cleaning_changes
    ]
    return pd.DataFrame(
        linhas,
        columns=[
            "linha",
            "coluna",
            "valor anterior",
            "valor posterior",
            "tipo da alteração",
            "motivo",
        ],
    )


def _indicadores(ws: Worksheet, indicadores: Sequence[Indicator]) -> None:
    """Uma tabela por indicador confirmado, com um gráfico de barras ao lado."""
    linha = 1
    for indicador in indicadores:
        ws.cell(row=linha, column=1, value=indicador.title).font = LABEL_FONT
        linha += 1
        ws.cell(row=linha, column=1, value=indicador.dimension).font = DATA_FONT
        ws.cell(row=linha, column=2, value=indicador.measure).font = DATA_FONT
        primeira_dado = linha + 1

        for chave, valor in indicador.rows:
            linha += 1
            ws.cell(row=linha, column=1, value=chave).font = DATA_FONT
            ws.cell(row=linha, column=2, value=round(valor, 2)).font = DATA_FONT

        if indicador.rows:
            grafico = BarChart()
            grafico.title = indicador.title
            grafico.height = 7
            grafico.width = 16
            ultima = min(linha, primeira_dado + _MAX_BARRAS - 1)
            grafico.add_data(
                Reference(ws, min_col=2, min_row=primeira_dado - 1, max_row=ultima),
                titles_from_data=True,
            )
            grafico.set_categories(Reference(ws, min_col=1, min_row=primeira_dado, max_row=ultima))
            ws.add_chart(grafico, f"E{primeira_dado}")

        linha += 1
        ws.cell(row=linha, column=1, value="Total").font = LABEL_FONT
        ws.cell(row=linha, column=2, value=round(indicador.total, 2)).font = LABEL_FONT
        if indicador.ignored_rows:
            linha += 1
            ws.cell(
                row=linha,
                column=1,
                value=(
                    f"{indicador.ignored_rows} linha(s) ignorada(s): o valor não pôde "
                    "ser lido como número"
                ),
            ).font = DATA_FONT
        linha += 3

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18


# ============================================================
# Entry point
# ============================================================


def write_analysis_report(path: Path, dados: ReportInput) -> Path:
    """
    Escreve o `relatorio_analise.xlsx`.

    Abas vazias são omitidas de propósito: uma aba "Antes e depois" em
    branco faria a pessoa procurar um erro que não existe.
    """
    workbook = Workbook()
    resumo = workbook.active
    if resumo is None:  # pragma: no cover
        resumo = workbook.create_sheet()
    resumo.title = _ABA_RESUMO
    _resumo(resumo, dados)

    opcionais: list[tuple[str, pd.DataFrame]] = [
        (_ABA_ABAS, _abas_frame(dados)),
        (_ABA_PROBLEMAS, _problemas_frame(dados)),
        (_ABA_REVISAO, _revisao_frame(dados)),
        (_ABA_ALTERACOES, _alteracoes_frame(dados)),
        (_ABA_ANTES_DEPOIS, _antes_depois_frame(dados)),
    ]
    for titulo, frame in opcionais:
        if frame.empty:
            continue
        write_dataframe_sheet(workbook.create_sheet(titulo), frame)

    if dados.indicators:
        _indicadores(workbook.create_sheet(_ABA_INDICADORES), dados.indicators)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


__all__ = ["ANALYSIS_REPORT_NAME", "ReportInput", "write_analysis_report"]
