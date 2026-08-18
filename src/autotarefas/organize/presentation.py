"""
Avaliacao OBJETIVA da apresentacao de uma planilha.

"Bonita" e "feia" nao sao criterios: duas pessoas discordam e ninguem
consegue conferir. Este modulo troca o julgamento por uma lista de
criterios verificaveis — cada um responde sim ou nao, com o detalhe que
justifica a resposta — e so entao emite um veredito.

Tres vereditos, e o que cada um significa para o produto:

    organizada   nao ha o que melhorar com seguranca; NAO oferecemos
                 reformatacao (mexer na identidade visual de quem ja
                 cuidou da planilha e desrespeito, nao ajuda)
    melhoravel   ha ganhos claros e seguros; oferecemos a organizacao
    ambigua      a estrutura nao esta clara o bastante para decidir
                 sozinho (celulas mescladas na area de dados, cabecalho
                 irreconhecivel) — quem decide e a pessoa

O modulo NAO altera nada: ele lê e conclui. Quem aplica e o `organizer`.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import TYPE_CHECKING, Any, Literal

from openpyxl import load_workbook

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.worksheet.worksheet import Worksheet

#: Veredito da avaliacao.
Verdict = Literal["organizada", "melhoravel", "ambigua"]

#: A partir de quantas linhas um filtro e um painel congelado passam a
#: fazer diferenca real para quem abre o arquivo.
LINHAS_PARA_FILTRO = 15
LINHAS_PARA_PAINEL = 15

#: Largura minima e maxima consideradas legiveis (unidades do Excel).
LARGURA_MINIMA = 6.0
LARGURA_MAXIMA = 80.0

#: Quantas cores de preenchimento distintas na area de dados ja indicam
#: excesso (o cabecalho nao entra na conta).
MAX_CORES_NOS_DADOS = 3

#: Amostra de linhas lidas para avaliar formato e alinhamento. Uma planilha
#: de 100 mil linhas nao precisa ser varrida inteira para se saber se a
#: coluna de datas tem formato consistente.
AMOSTRA = 200

#: Menos de duas colunas preenchidas nao formam cabecalho de tabela.
MIN_TITULOS = 2

#: Uma linha solta abaixo da tabela e nota de rodape; duas ja sao um corpo.
MIN_LINHAS_DE_OUTRA_TABELA = 2

#: Criterios ESTRUTURAIS: falhar num deles significa que nao entendemos a
#: tabela, e organizar seria chutar. Falta de destaque no cabecalho NAO
#: entra aqui — isso e questao de apresentacao, e tem conserto seguro.
_ESTRUTURAIS = frozenset({"estrutura_tabular", "cabecalho_presente", "tabela_unica"})


@dataclass(frozen=True, slots=True)
class Criterion:
    """Um criterio verificavel — com a evidencia que sustenta o resultado."""

    key: str
    title: str
    passed: bool
    detail: str
    applicable: bool = True


@dataclass(frozen=True, slots=True)
class PresentationAudit:
    """O que a planilha ja tem, o que falta, e o veredito."""

    sheet: str
    verdict: Verdict
    criteria: tuple[Criterion, ...] = ()
    rows: int = 0
    columns: int = 0

    @property
    def applicable(self) -> tuple[Criterion, ...]:
        return tuple(c for c in self.criteria if c.applicable)

    @property
    def failures(self) -> tuple[Criterion, ...]:
        return tuple(c for c in self.applicable if not c.passed)

    @property
    def score(self) -> float:
        """Fracao dos criterios aplicaveis atendidos (0..1)."""
        aplicaveis = self.applicable
        if not aplicaveis:
            return 1.0
        return sum(1 for c in aplicaveis if c.passed) / len(aplicaveis)

    @property
    def can_improve(self) -> bool:
        """Vale oferecer a organizacao?"""
        return self.verdict == "melhoravel"

    def as_dict(self) -> dict[str, Any]:
        return {
            "aba": self.sheet,
            "veredito": self.verdict,
            "pontuacao": round(self.score, 3),
            "linhas": self.rows,
            "colunas": self.columns,
            "criterios": [
                {
                    "chave": c.key,
                    "titulo": c.title,
                    "atendido": c.passed,
                    "detalhe": c.detail,
                    "aplicavel": c.applicable,
                }
                for c in self.criteria
            ],
            # Com o titulo so, a tela dizia "Estrutura tabular clara" e a
            # pessoa tinha de abrir o relatorio para saber POR QUE falhou.
            "pendencias": [f"{c.title} — {c.detail}" for c in self.failures],
        }


# ============================================================
# Leitura de fatos (nada de julgamento aqui)
# ============================================================


def _largura(ws: Worksheet, indice: int) -> float | None:
    from openpyxl.utils import get_column_letter

    dimensao = ws.column_dimensions.get(get_column_letter(indice))
    return None if dimensao is None or dimensao.width is None else float(dimensao.width)


def _celulas_de_dados(ws: Worksheet, header_row: int) -> list[list[Any]]:
    """Amostra das celulas da area de dados (objetos, nao valores)."""
    linhas: list[list[Any]] = []
    for linha in ws.iter_rows(min_row=header_row + 1, max_row=header_row + AMOSTRA):
        if all(c.value is None for c in linha):
            continue
        linhas.append(list(linha))
    return linhas


def _mescladas_na_area(ws: Worksheet, header_row: int) -> int:
    """Celulas mescladas DENTRO da area de dados (as piores para ler)."""
    return sum(
        1
        for intervalo in ws.merged_cells.ranges
        if intervalo.max_row >= header_row  # inclui o proprio cabecalho
    )


def _cores_dos_dados(amostra: list[list[Any]]) -> int:
    cores: set[str] = set()
    for linha in amostra:
        for celula in linha:
            preenchimento = celula.fill
            if preenchimento is not None and preenchimento.fill_type == "solid":
                cores.add(str(preenchimento.fgColor.rgb))
    # Branco/sem cor nao conta como "cor".
    return len({c for c in cores if c not in {"00000000", "FFFFFFFF", "None"}})


# ============================================================
# Criterios
# ============================================================


def _criterios_cabecalho(ws: Worksheet, header_row: int) -> tuple[Criterion, Criterion]:
    """
    Duas perguntas diferentes sobre a mesma linha.

    "Existe cabecalho?" e ESTRUTURAL: sem titulos nao ha tabela para
    organizar. "Ele se destaca?" e APRESENTACAO: tem conserto seguro, e
    por isso nao pode empurrar a planilha para o veredito 'ambigua'.
    """
    celulas = [c for c in ws[header_row] if c.value is not None]
    presente = Criterion(
        "cabecalho_presente",
        "Cabeçalho preenchido",
        passed=bool(celulas),
        detail=(
            f"{len(celulas)} título(s) na linha {header_row}"
            if celulas
            else "a linha de cabeçalho não tem nenhum título preenchido"
        ),
    )

    def destacada(celula: Any) -> bool:
        preenchimento = celula.fill
        colorida = (
            preenchimento is not None
            and preenchimento.fill_type == "solid"
            and str(preenchimento.fgColor.rgb) not in {"00000000", "FFFFFFFF"}
        )
        return bool(celula.font.bold) or colorida

    destaque = Criterion(
        "cabecalho_destacado",
        "Cabeçalho legível e destacado",
        passed=any(destacada(c) for c in celulas),
        detail=(
            "títulos com negrito ou preenchimento"
            if any(destacada(c) for c in celulas)
            else "títulos sem destaque visual (nem negrito, nem preenchimento)"
        ),
        applicable=bool(celulas),
    )
    return presente, destaque


def _criterio_larguras(ws: Worksheet, header_row: int, colunas: int) -> Criterion:
    """
    Largura da coluna e titulo visivel — um criterio so.

    Eram dois, e o relatorio saia contraditorio: "Larguras legiveis: OK" na
    linha de cima e "Titulos visiveis por inteiro: melhorar" na de baixo. Para
    quem le, largura e o titulo caber sao a mesma pergunta. Continuam duas
    medicoes, com um veredito so.
    """
    fora_de_faixa: list[str] = []
    cortados: list[str] = []
    for indice in range(1, colunas + 1):
        largura = _largura(ws, indice)
        if largura is None:
            continue
        titulo = str(ws.cell(row=header_row, column=indice).value or "")
        rotulo = titulo or f"coluna {indice}"
        if largura < LARGURA_MINIMA or largura > LARGURA_MAXIMA:
            fora_de_faixa.append(rotulo)
        elif titulo and len(titulo) > largura:
            cortados.append(rotulo)

    if not fora_de_faixa and not cortados:
        detalhe = "larguras em faixa legível e nenhum título cortado"
    elif fora_de_faixa and cortados:
        detalhe = (
            f"fora da faixa: {', '.join(fora_de_faixa[:3])}; cortado(s): {', '.join(cortados[:3])}"
        )
    elif cortados:
        detalhe = f"título(s) provavelmente cortado(s): {', '.join(cortados[:4])}"
    else:
        detalhe = f"largura fora da faixa legível: {', '.join(fora_de_faixa[:4])}"

    return Criterion(
        "larguras_legiveis",
        "Larguras e títulos legíveis",
        passed=not (fora_de_faixa or cortados),
        detail=detalhe,
    )


def _criterio_formatos(amostra: list[list[Any]], colunas: int) -> Criterion:
    inconsistentes: list[int] = []
    for indice in range(colunas):
        formatos = {
            linha[indice].number_format
            for linha in amostra
            if indice < len(linha) and linha[indice].value is not None
        }
        if len(formatos) > 1:
            inconsistentes.append(indice + 1)
    return Criterion(
        "formatos_consistentes",
        "Formatos consistentes por coluna",
        passed=not inconsistentes,
        detail=(
            "cada coluna usa um formato só"
            if not inconsistentes
            else f"{len(inconsistentes)} coluna(s) com formatos misturados"
        ),
        applicable=bool(amostra),
    )


def _criterio_alinhamento(amostra: list[list[Any]], colunas: int) -> Criterion:
    misturadas = 0
    for indice in range(colunas):
        alinhamentos = {
            linha[indice].alignment.horizontal
            for linha in amostra
            if indice < len(linha) and linha[indice].value is not None
        }
        if len({a for a in alinhamentos if a}) > 1:
            misturadas += 1
    return Criterion(
        "alinhamento_coerente",
        "Alinhamento coerente",
        passed=misturadas == 0,
        detail=(
            "alinhamento uniforme dentro de cada coluna"
            if misturadas == 0
            else f"{misturadas} coluna(s) com alinhamentos diferentes"
        ),
        applicable=bool(amostra),
    )


def _criterio_filtro(ws: Worksheet, linhas: int) -> Criterion:
    tem = ws.auto_filter.ref is not None or bool(getattr(ws, "tables", {}))
    aplicavel = linhas >= LINHAS_PARA_FILTRO
    return Criterion(
        "filtro",
        "Filtro na área tabular",
        passed=tem,
        detail="filtro presente" if tem else f"sem filtro (a planilha tem {linhas} linhas)",
        applicable=aplicavel,
    )


def _criterio_painel(ws: Worksheet, linhas: int) -> Criterion:
    tem = bool(ws.freeze_panes)
    aplicavel = linhas >= LINHAS_PARA_PAINEL
    return Criterion(
        "painel_congelado",
        "Cabeçalho congelado",
        passed=tem,
        detail=(
            f"painel congelado em {ws.freeze_panes}"
            if tem
            else f"sem congelamento (a planilha tem {linhas} linhas)"
        ),
        applicable=aplicavel,
    )


def _criterio_cores(amostra: list[list[Any]]) -> Criterion:
    cores = _cores_dos_dados(amostra)
    return Criterion(
        "cores_moderadas",
        "Uso moderado de cores",
        passed=cores <= MAX_CORES_NOS_DADOS,
        detail=f"{cores} cor(es) de preenchimento na área de dados",
        applicable=bool(amostra),
    )


def _criterio_estrutura(ws: Worksheet, header_row: int, colunas: int) -> Criterion:
    mescladas = _mescladas_na_area(ws, header_row)
    titulos = [str(c.value).strip() for c in ws[header_row] if c.value is not None]
    duplicados = len(titulos) != len(set(titulos))
    vazios = colunas - len(titulos)

    problemas: list[str] = []
    if mescladas:
        problemas.append(f"{mescladas} célula(s) mesclada(s) na área de dados")
    if duplicados:
        problemas.append("títulos de coluna repetidos")
    if vazios > 0:
        problemas.append(f"{vazios} coluna(s) sem título")

    return Criterion(
        "estrutura_tabular",
        "Estrutura tabular clara",
        passed=not problemas,
        detail="; ".join(problemas) if problemas else "uma tabela simples, sem mesclagens",
    )


#: Linha de dados: o numero fisico e os valores preenchidos dela.
_LinhaDeDados = tuple[int, list[Any]]


def _blocos_de_dados(ws: Worksheet, header_row: int) -> list[list[_LinhaDeDados]]:
    """
    Agrupa as linhas de dados em blocos separados por linhas vazias.

    Uma passada so, com `values_only`. A versao anterior pedia a linha inteira
    ao openpyxl uma por vez (`ws[n]`), o que construia objetos de celula a cada
    volta: numa planilha de 7 mil linhas isso custava 16 segundos e fazia a
    analise inteira parecer travada.
    """
    blocos: list[list[_LinhaDeDados]] = []
    atual: list[_LinhaDeDados] = []
    for deslocamento, valores in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True)):
        preenchidas = [v for v in valores if v is not None and str(v).strip() != ""]
        if not preenchidas:
            if atual:
                blocos.append(atual)
                atual = []
            continue
        atual.append((header_row + 1 + deslocamento, preenchidas))
    if atual:
        blocos.append(atual)
    return blocos


def _criterio_tabela_unica(ws: Worksheet, header_row: int) -> Criterion:
    """
    Ha mais de uma tabela empilhada na mesma aba?

    Duas bases coladas com uma linha em branco no meio sao lidas como uma
    tabela so, e a contagem de registros sai errada sem ninguem perceber. Nao
    da para adivinhar qual delas a pessoa quis: viramos a pergunta para ela.

    O sinal e conservador de proposito — linha em branco no meio de uma base e
    comum demais para virar alarme sozinha. Exige-se, alem do intervalo: um
    bloco seguinte com corpo (2+ linhas), a primeira linha desse bloco toda
    em texto (cara de cabecalho) e uma forma diferente da do primeiro bloco.
    """
    blocos = _blocos_de_dados(ws, header_row)
    suspeitos: list[int] = []
    if len(blocos) > 1:
        largura_inicial = len(blocos[0][0][1])
        tipos_iniciais = {type(v) for _, valores in blocos[0][:AMOSTRA] for v in valores}
        for bloco in blocos[1:]:
            if len(bloco) < MIN_LINHAS_DE_OUTRA_TABELA:  # nota solta, nao tabela
                continue
            linha, valores = bloco[0]
            if len(valores) < MIN_TITULOS or not all(isinstance(v, str) for v in valores):
                continue
            # Mesma forma e mesmo feitio dos dados: e continuacao, nao tabela nova.
            if len(valores) == largura_inicial and tipos_iniciais <= {str}:
                continue
            suspeitos.append(linha)

    return Criterion(
        "tabela_unica",
        "Uma tabela por aba",
        passed=not suspeitos,
        detail=(
            "os dados formam uma tabela contínua"
            if not suspeitos
            else (
                "parece haver outra tabela na mesma aba, começando na linha "
                f"{suspeitos[0]}: a contagem de registros mistura as duas. "
                "Separe em abas diferentes ou escolha uma delas"
            )
        ),
    )


def _criterio_tabela(ws: Worksheet) -> Criterion:
    tabelas = getattr(ws, "tables", {}) or {}
    return Criterion(
        "tabela_estruturada",
        "Tabela estruturada do Excel",
        passed=bool(tabelas),
        detail=f"{len(tabelas)} tabela(s) declarada(s)" if tabelas else "sem tabela declarada",
        # Nao e obrigatorio: uma planilha pode estar impecavel sem isso.
        applicable=False,
    )


# ============================================================
# Veredito
# ============================================================


def _veredito(criterios: tuple[Criterion, ...]) -> Verdict:
    """
    Do conjunto de criterios para uma decisao de produto.

    Falha ESTRUTURAL (mesclagem na area de dados, cabecalho irreconhecivel)
    nao e questao de gosto: sem entender a tabela, organizar seria chutar.
    Nesse caso o veredito e `ambigua` e quem decide e a pessoa.
    """
    aplicaveis = [c for c in criterios if c.applicable]
    falhas = [c for c in aplicaveis if not c.passed]
    if any(c.key in _ESTRUTURAIS for c in falhas):
        return "ambigua"
    return "organizada" if not falhas else "melhoravel"


def audit_presentation(
    path: Path,
    *,
    sheet: str | None = None,
    header_row: int = 1,
) -> PresentationAudit:
    """
    Avalia a apresentacao de uma aba, criterio a criterio.

    Args:
        path: arquivo XLSX/XLSM (nunca e alterado).
        sheet: aba a avaliar. None = a ativa.
        header_row: linha fisica do cabecalho.

    Returns:
        PresentationAudit com o veredito e a evidencia de cada criterio.
    """
    workbook = load_workbook(path, read_only=False, data_only=False)
    ws = workbook[sheet] if sheet is not None else workbook.active
    if ws is None:  # pragma: no cover - openpyxl sempre devolve uma aba
        msg = "planilha sem aba ativa"
        raise ValueError(msg)

    colunas = int(ws.max_column or 0)
    linhas_dados = max(int(ws.max_row or 0) - header_row, 0)
    amostra = _celulas_de_dados(ws, header_row)

    cabecalho_presente, cabecalho_destacado = _criterios_cabecalho(ws, header_row)
    criterios = (
        cabecalho_presente,
        cabecalho_destacado,
        _criterio_estrutura(ws, header_row, colunas),
        _criterio_tabela_unica(ws, header_row),
        _criterio_larguras(ws, header_row, colunas),
        _criterio_formatos(amostra, colunas),
        _criterio_alinhamento(amostra, colunas),
        _criterio_filtro(ws, linhas_dados),
        _criterio_painel(ws, linhas_dados),
        _criterio_cores(amostra),
        _criterio_tabela(ws),
    )

    workbook.close()
    return PresentationAudit(
        sheet=str(ws.title),
        verdict=_veredito(criterios),
        criteria=criterios,
        rows=linhas_dados,
        columns=colunas,
    )


#: Trechos entre colchetes sao locale e cor ("[$-409]", "[Red]"), nao data.
_COLCHETES = re.compile(r"\[[^\]]*\]")


def _mes_antes_do_dia(formato: str) -> bool:
    """
    O formato mostra o mes ANTES do dia (padrao americano)?

    Comparar a ordem de aparicao de `y`, `m` e `d` distingue os tres casos que
    importam sem depender de listar formatos um a um:

        mm-dd-yy    -> m antes de d  -> americano
        dd/mm/aaaa  -> d antes de m  -> brasileiro
        yyyy-mm-dd  -> y primeiro    -> ISO, tambem sem ambiguidade
    """
    limpo = _COLCHETES.sub("", formato).lower()
    posicoes = {letra: limpo.find(letra) for letra in "ymd"}
    if posicoes["d"] < 0 or posicoes["m"] < 0:
        return False  # sem dia e mes juntos nao e data (ex.: "h:mm" e minuto)
    if 0 <= posicoes["y"] < posicoes["m"]:
        return False  # ano primeiro: ISO
    return posicoes["m"] < posicoes["d"]


#: Fracao da coluna que precisa parecer numero para valer o aviso.
_MAIORIA_NUMERICA = 0.6


def _parece_numero(texto: str) -> bool:
    """
    O texto e um numero escrito por gente? (R$ 1.234,50 / 15% / -3,5)

    Identificadores com zero a esquerda ficam de fora: "000123" e codigo, nao
    quantidade, e converte-lo seria justamente o estrago que o card evita.
    """
    limpo = texto.strip().replace("R$", "").replace("%", "").strip()
    if not limpo:
        return False
    negativo = limpo.startswith("-")
    corpo = limpo[1:] if negativo else limpo
    if len(corpo) > 1 and corpo[0] == "0" and corpo[1].isdigit():
        return False  # zero a esquerda: identificador
    corpo = corpo.replace(".", "").replace(",", ".")
    try:
        float(corpo)
    except ValueError:
        return False
    return True


def text_number_notes(
    path: Path, *, sheet: str | None = None, header_row: int = 1
) -> tuple[str, ...]:
    """
    Observa colunas de numeros guardados como TEXTO — sem converter nada.

    O leitor entende o valor de qualquer jeito, entao a analise nao quebra. Mas
    dentro do Excel essa coluna nao soma, nao ordena direito e nao entra em
    formula: e um problema real da planilha, e quem decide corrigir e o dono.

    Args:
        path: XLSX/XLSM. Aberto somente para leitura.
        sheet: aba analisada; `None` usa a ativa.
        header_row: linha dos titulos.

    Returns:
        Uma observacao por coluna encontrada. Vazio quando nao ha nenhuma.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = workbook[sheet] if sheet and sheet in workbook.sheetnames else workbook.active
        if ws is None:  # pragma: no cover - arquivo sem aba ativa
            return ()
        achadas: list[str] = []
        for indice in range(1, int(ws.max_column or 0) + 1):
            valores = [
                linha[0].value
                for linha in ws.iter_rows(
                    min_row=header_row + 1,
                    max_row=header_row + AMOSTRA,
                    min_col=indice,
                    max_col=indice,
                )
                if linha and linha[0].value is not None
            ]
            textos = [v for v in valores if isinstance(v, str) and v.strip()]
            if not valores or len(textos) < len(valores) * _MAIORIA_NUMERICA:
                continue
            numericos = [v for v in textos if _parece_numero(v)]
            if len(numericos) < len(textos) * _MAIORIA_NUMERICA:
                continue
            titulo = str(ws.cell(row=header_row, column=indice).value or f"coluna {indice}")
            achadas.append(
                f"a coluna '{titulo}' guarda números como texto (ex.: "
                f"'{numericos[0]}'). Dentro do Excel eles não somam nem ordenam "
                "como número. Nada foi convertido — a decisão é sua"
            )
        return tuple(achadas)
    finally:
        workbook.close()


def date_format_notes(
    path: Path, *, sheet: str | None = None, header_row: int = 1
) -> tuple[str, ...]:
    """
    Observa colunas de data em formato americano — sem alterar NADA.

    Preservar o formato do arquivo continua sendo a regra. Mas ficar calado
    quando uma planilha brasileira mostra 12-01-19 no lugar de 01/12/2019 e
    deixar a pessoa descobrir sozinha, e ela tem o direito de saber.

    Args:
        path: XLSX/XLSM. Aberto somente para leitura.
        sheet: aba analisada; `None` usa a ativa.
        header_row: linha dos titulos.

    Returns:
        Uma observacao por coluna encontrada. Vazio quando nao ha nenhuma.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = workbook[sheet] if sheet and sheet in workbook.sheetnames else workbook.active
        if ws is None:  # pragma: no cover - arquivo sem aba ativa
            return ()
        colunas = int(ws.max_column or 0)
        achadas: list[str] = []
        for indice in range(1, colunas + 1):
            celulas = [
                linha[0]
                for linha in ws.iter_rows(
                    min_row=header_row + 1,
                    max_row=header_row + AMOSTRA,
                    min_col=indice,
                    max_col=indice,
                )
                if linha
            ]
            # A barra invertida e so escape do Excel ("mm\-dd\-yy"); tira-la
            # evita ter que adivinhar onde ela cai dentro do formato.
            formatos = {
                str(c.number_format).replace("\\", "") for c in celulas if c.value is not None
            }
            americanos = sorted(f for f in formatos if _mes_antes_do_dia(f))
            if not americanos:
                continue
            titulo = str(ws.cell(row=header_row, column=indice).value or f"coluna {indice}")
            achadas.append(
                f"a coluna '{titulo}' exibe datas no formato americano "
                f"({americanos[0]}), com o mês antes do dia. Os valores não foram "
                "alterados — trocar o formato mudaria como a planilha é lida"
            )
        return tuple(achadas)
    finally:
        workbook.close()


__all__ = [
    "AMOSTRA",
    "LARGURA_MAXIMA",
    "LARGURA_MINIMA",
    "LINHAS_PARA_FILTRO",
    "LINHAS_PARA_PAINEL",
    "MAX_CORES_NOS_DADOS",
    "MIN_LINHAS_DE_OUTRA_TABELA",
    "MIN_TITULOS",
    "Criterion",
    "PresentationAudit",
    "Verdict",
    "audit_presentation",
    "date_format_notes",
    "text_number_notes",
]
