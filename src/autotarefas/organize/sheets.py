"""
Inspecao das ABAS de uma pasta de trabalho.

Uma planilha empresarial raramente tem uma aba so: costuma ter a base, uma
capa, um "aux" com listas de validacao e duas ou tres esquecidas vazias.
Formatar tudo indiscriminadamente seria estragar a capa; escolher a
primeira em silencio seria processar a aba errada.

Este modulo classifica cada aba em uma de cinco naturezas e devolve o que
a interface precisa para PERGUNTAR quando houver mais de uma candidata:

    dados        parece uma tabela: cabecalho reconhecivel + linhas abaixo
    apresentacao poucas celulas, muito texto solto (capa, instrucoes)
    auxiliar     listas curtas de apoio (validacao, de/para)
    vazia        nenhuma celula preenchida
    ambigua      tem conteudo, mas nao da para afirmar o que e

Nao decide nada: descreve. Quem escolhe e a pessoa.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import TYPE_CHECKING, Any, Literal

from openpyxl import load_workbook

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.worksheet.worksheet import Worksheet

#: Natureza de uma aba.
SheetKind = Literal["dados", "apresentacao", "auxiliar", "vazia", "ambigua"]

#: Uma tabela de verdade tem pelo menos duas colunas e uma linha de dados.
MIN_COLUNAS_TABULAR = 2
MIN_LINHAS_TABULAR = 1

#: Acima disto a aba deixa de ser "auxiliar" e vira base de dados.
MAX_LINHAS_AUXILIAR = 10

#: Densidade minima (celulas preenchidas / celulas da AMOSTRA) para uma aba
#: parecer tabela. Uma capa tem densidade baixissima.
DENSIDADE_TABULAR = 0.45

#: Quantas linhas do topo bastam para classificar a aba. Ler 200 mil linhas
#: so para dizer "isto e uma tabela" seria caro e inutil.
LINHAS_AMOSTRADAS = 50


@dataclass(frozen=True, slots=True)
class SheetInfo:
    """O que se sabe sobre uma aba — e por que ela foi classificada assim."""

    name: str
    kind: SheetKind
    rows: int
    columns: int
    filled_cells: int
    reason: str
    header_row: int | None = None

    @property
    def is_candidate(self) -> bool:
        """Aba que faz sentido processar."""
        return self.kind == "dados"

    def as_dict(self) -> dict[str, Any]:
        return {
            "nome": self.name,
            "natureza": self.kind,
            "linhas": self.rows,
            "colunas": self.columns,
            "celulas_preenchidas": self.filled_cells,
            "motivo": self.reason,
            "linha_do_cabecalho": self.header_row,
            "candidata": self.is_candidate,
        }


def _fatos(ws: Worksheet) -> tuple[int, int, int, int | None, int]:
    """
    (linhas, colunas, celulas preenchidas, cabecalho, linhas amostradas).

    As celulas sao contadas so nas primeiras `LINHAS_AMOSTRADAS` linhas. Quem
    usar esse numero PRECISA dividir pela area da amostra, nunca pela area da
    planilha inteira — foi exatamente esse descompasso que fazia toda tabela
    com mais de ~111 linhas ser classificada como "ambigua".
    """
    linhas = int(ws.max_row or 0)
    colunas = int(ws.max_column or 0)
    amostradas = min(linhas, LINHAS_AMOSTRADAS)

    preenchidas = 0
    cabecalho: int | None = None
    for indice, linha in enumerate(ws.iter_rows(max_row=amostradas), start=1):
        valores = [c.value for c in linha if c.value is not None]
        preenchidas += len(valores)
        # A primeira linha com >= 2 celulas preenchidas e a candidata natural
        # a cabecalho — a mesma heuristica que o leitor usa.
        if cabecalho is None and len(valores) >= MIN_COLUNAS_TABULAR:
            cabecalho = indice
    return linhas, colunas, preenchidas, cabecalho, amostradas


def _classificar(
    linhas: int, colunas: int, preenchidas: int, cabecalho: int | None, amostradas: int
) -> tuple[SheetKind, str]:
    if preenchidas == 0:
        return "vazia", "nenhuma célula preenchida"

    if colunas < MIN_COLUNAS_TABULAR or cabecalho is None:
        return "apresentacao", "poucas colunas preenchidas: parece capa ou texto solto"

    linhas_de_dados = max(linhas - cabecalho, 0)
    if linhas_de_dados < MIN_LINHAS_TABULAR:
        return "apresentacao", "há um cabeçalho, mas nenhuma linha de dados abaixo"

    # A area e a DA AMOSTRA, porque `preenchidas` so contou a amostra.
    area = max(amostradas * colunas, 1)
    densidade = preenchidas / area
    if densidade < DENSIDADE_TABULAR:
        return (
            "ambigua",
            f"área muito esparsa ({densidade:.0%} preenchida) para afirmar que é uma tabela",
        )

    if linhas_de_dados <= MAX_LINHAS_AUXILIAR and colunas <= MIN_COLUNAS_TABULAR:
        return "auxiliar", f"lista curta de apoio ({linhas_de_dados} linha(s), {colunas} coluna(s))"

    return (
        "dados",
        f"tabela com cabeçalho na linha {cabecalho} e {linhas_de_dados} linha(s) de dados",
    )


def survey_sheets(path: Path) -> tuple[SheetInfo, ...]:
    """
    Inspeciona TODAS as abas do arquivo e classifica cada uma.

    Args:
        path: arquivo XLSX/XLSM (nunca alterado). Para CSV, quem chama nao
            precisa desta funcao: ha uma "aba" so.

    Returns:
        Uma `SheetInfo` por aba, na ordem do arquivo.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        resultado: list[SheetInfo] = []
        for nome in workbook.sheetnames:
            ws = workbook[nome]
            linhas, colunas, preenchidas, cabecalho, amostradas = _fatos(ws)
            natureza, motivo = _classificar(linhas, colunas, preenchidas, cabecalho, amostradas)
            resultado.append(
                SheetInfo(
                    name=str(nome),
                    kind=natureza,
                    rows=linhas,
                    columns=colunas,
                    filled_cells=preenchidas,
                    reason=motivo,
                    header_row=cabecalho,
                )
            )
        return tuple(resultado)
    finally:
        workbook.close()


def candidates(sheets: tuple[SheetInfo, ...]) -> tuple[SheetInfo, ...]:
    """As abas que fazem sentido processar."""
    return tuple(s for s in sheets if s.is_candidate)


def needs_sheet_choice(sheets: tuple[SheetInfo, ...]) -> bool:
    """Ha mais de uma candidata? Entao a escolha e da pessoa, nao nossa."""
    return len(candidates(sheets)) > 1


__all__ = [
    "DENSIDADE_TABULAR",
    "LINHAS_AMOSTRADAS",
    "MAX_LINHAS_AUXILIAR",
    "MIN_COLUNAS_TABULAR",
    "MIN_LINHAS_TABULAR",
    "SheetInfo",
    "SheetKind",
    "candidates",
    "needs_sheet_choice",
    "survey_sheets",
]
