"""
Apresentacao comum dos artefatos XLSX do AutoTarefas.

Paleta, fontes e as quatro operacoes que todo artefato de planilha do
projeto repete: converter uma celula para um tipo que o openpyxl aceite,
formatar o cabecalho (com painel congelado e autofiltro), ajustar a
largura das colunas e despejar um DataFrame numa aba.

Existe para que `planilha_validada.xlsx` (PLA-007) e `divergencias.xlsx`
(REC-001) tenham a MESMA cara para o cliente — e para que ajustar essa
cara seja uma mudanca em um lugar so.

Isto e apresentacao: nenhuma funcao daqui altera o significado de um
valor. O que entra e o que sai.
"""

from __future__ import annotations

import math
from typing import TYPE_CHECKING

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    import pandas as pd
    from openpyxl.worksheet.worksheet import Worksheet

#: Fonte usada em todos os artefatos (presente em Excel e LibreOffice).
FONT_NAME = "Arial"

TITLE_FONT = Font(name=FONT_NAME, bold=True, size=16, color="1F4E78")
LABEL_FONT = Font(name=FONT_NAME, bold=True)
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
DATA_FONT = Font(name=FONT_NAME)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")  # azul escuro
OK_FILL = PatternFill("solid", fgColor="E2EFDA")  # verde claro
ERROR_FILL = PatternFill("solid", fgColor="FCE4E4")  # vermelho claro
NEUTRAL_FILL = PatternFill("solid", fgColor="FFF2CC")  # amarelo claro

CENTER = Alignment(horizontal="center", vertical="center")

#: Teto de largura de coluna: sem ele, uma celula longa empurra o resto da
#: planilha para fora da tela.
MAX_COLUMN_WIDTH = 60


def cell_value(value: object) -> object:
    """Converte uma celula do DataFrame para um tipo que o openpyxl aceita."""
    # None ou NaN viram string vazia.
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    if isinstance(value, (bool, int, float, str)):
        return value
    # numpy int64/float64 e afins: converte para nativo via item().
    item = getattr(value, "item", None)
    if callable(item):
        native = item()
        if isinstance(native, (bool, int, float, str)):
            return native
    return str(value)


def style_header_row(ws: Worksheet, n_cols: int) -> None:
    """Formata o cabecalho (linha 1), congela o painel e liga o autofiltro."""
    if n_cols <= 0:
        return
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
    ws.freeze_panes = "A2"
    last_col = get_column_letter(n_cols)
    ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"


def autofit_columns(ws: Worksheet, max_width: int = MAX_COLUMN_WIDTH) -> None:
    """Ajusta a largura das colunas ao maior conteudo (com um teto)."""
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[letter].width = min(max(length + 2, 10), max_width)


def write_dataframe_sheet(ws: Worksheet, dataframe: pd.DataFrame) -> None:
    """Escreve cabecalho + linhas do DataFrame e aplica a formatacao padrao."""
    ws.append([str(c) for c in dataframe.columns])
    for row in dataframe.itertuples(index=False, name=None):
        ws.append([cell_value(v) for v in row])
    for cell_row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in cell_row:
            cell.font = DATA_FONT
    style_header_row(ws, len(dataframe.columns))
    autofit_columns(ws)


__all__ = [
    "CENTER",
    "DATA_FONT",
    "ERROR_FILL",
    "FONT_NAME",
    "HEADER_FILL",
    "HEADER_FONT",
    "LABEL_FONT",
    "MAX_COLUMN_WIDTH",
    "NEUTRAL_FILL",
    "OK_FILL",
    "TITLE_FONT",
    "autofit_columns",
    "cell_value",
    "style_header_row",
    "write_dataframe_sheet",
]
