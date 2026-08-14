"""
Testes do comando `autotarefas corrigir` (RF-PLA-010), de ponta a ponta.

Um XLSX real + um arquivo de regras confirmadas entram; a planilha
corrigida (com a apresentacao do original preservada), os itens para
revisao e o relatorio das quatro naturezas saem.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pandas as pd
import pytest
from click.testing import CliRunner
from openpyxl import load_workbook

from autotarefas.cli.commands.corrigir import corrigir
from autotarefas.cli.context import CLIContext
from autotarefas.tasks.correction_artifacts import (
    CORRECTED_CSV_NAME,
    CORRECTED_XLSX_NAME,
    JSON_REPORT_NAME,
    REVIEW_CSV_NAME,
)

FIXTURES = Path(__file__).parent.parent / "fixtures" / "correcoes"
BASE = FIXTURES / "base_para_corrigir.xlsx"
REGRAS = FIXTURES / "regras.yaml"


@pytest.fixture
def cli_ctx() -> CLIContext:
    return CLIContext()


def _run(ctx: CLIContext, *extra: str):
    return CliRunner().invoke(corrigir, [str(BASE), "--regras", str(REGRAS), *extra], obj=ctx)


class TestExecucao:
    def test_exit_0_e_resumo(self, cli_ctx: CLIContext) -> None:
        result = _run(cli_ctx)
        assert result.exit_code == 0
        assert "Regras confirmadas aplicadas: 3" in result.output
        assert "'Sao Paulo'  ->  'SP'" in result.output

    def test_item_fora_da_regra_aparece_no_terminal(self, cli_ctx: CLIContext) -> None:
        result = _run(cli_ctx)
        assert "Sao Jorge" in result.output
        assert "Para revisao" in result.output

    def test_regras_invalidas_sao_erro_de_uso(self, cli_ctx: CLIContext, tmp_path: Path) -> None:
        ruim = tmp_path / "regras.yaml"
        ruim.write_text("correcoes: []", encoding="utf-8")
        result = CliRunner().invoke(corrigir, [str(BASE), "--regras", str(ruim)], obj=cli_ctx)
        assert result.exit_code == 2

    def test_falhar_se_revisao(self, cli_ctx: CLIContext) -> None:
        result = _run(cli_ctx, "--falhar-se-revisao")
        assert result.exit_code == 1


class TestArtefatos:
    def test_gera_os_quatro_artefatos(self, cli_ctx: CLIContext, tmp_path: Path) -> None:
        result = _run(cli_ctx, "--out-dir", str(tmp_path))
        assert result.exit_code == 0
        for nome in (
            CORRECTED_XLSX_NAME,
            CORRECTED_CSV_NAME,
            REVIEW_CSV_NAME,
            JSON_REPORT_NAME,
        ):
            assert (tmp_path / nome).exists(), nome

    def test_planilha_corrigida_tem_os_valores_canonicos(
        self, cli_ctx: CLIContext, tmp_path: Path
    ) -> None:
        _run(cli_ctx, "--out-dir", str(tmp_path))
        base = pd.read_csv(tmp_path / CORRECTED_CSV_NAME, dtype=str).fillna("")

        assert base.loc[base["codigo"] == "00123", "uf"].item() == "SP"
        assert base.loc[base["codigo"] == "00124", "situacao"].item() == "Inativo"
        assert base.loc[base["codigo"] == "00123", "origem"].item() == "planilha"

    def test_valor_fora_da_regra_nao_e_alterado(self, cli_ctx: CLIContext, tmp_path: Path) -> None:
        _run(cli_ctx, "--out-dir", str(tmp_path))
        base = pd.read_csv(tmp_path / CORRECTED_CSV_NAME, dtype=str).fillna("")
        assert base.loc[base["codigo"] == "00126", "uf"].item() == "Sao Jorge"

    def test_xlsx_corrigido_mantem_os_dados_do_original(
        self, cli_ctx: CLIContext, tmp_path: Path
    ) -> None:
        _run(cli_ctx, "--out-dir", str(tmp_path))
        ws = load_workbook(tmp_path / CORRECTED_XLSX_NAME)["Dados"]

        assert [c.value for c in ws[1]] == ["codigo", "cliente", "uf", "situacao", "origem"]
        assert ws["C2"].value == "SP"
        assert ws["B2"].value == "Ana Lima"  # coluna sem regra: intacta

    def test_itens_para_revisao(self, cli_ctx: CLIContext, tmp_path: Path) -> None:
        _run(cli_ctx, "--out-dir", str(tmp_path))
        revisao = pd.read_csv(tmp_path / REVIEW_CSV_NAME, dtype=str)
        assert list(revisao["valor"]) == ["Sao Jorge"]

    def test_relatorio_separa_as_naturezas(self, cli_ctx: CLIContext, tmp_path: Path) -> None:
        _run(cli_ctx, "--out-dir", str(tmp_path))
        payload = json.loads((tmp_path / JSON_REPORT_NAME).read_text(encoding="utf-8"))

        assert payload["requisito"] == "RF-PLA-010"
        assert set(payload["resumo"]) == {
            "automatico_seguro",
            "normalizacao",
            "regra_confirmada",
            "revisao",
        }
        assert payload["resumo"]["regra_confirmada"] == 7
        assert payload["resumo"]["revisao"] == 1
        assert payload["limitacoes"]

    def test_dry_run_nao_grava(self, tmp_path: Path) -> None:
        destino = tmp_path / "saida"
        result = _run(CLIContext(dry_run=True), "--out-dir", str(destino))
        assert result.exit_code == 0
        assert "DRY-RUN" in result.output
        assert not destino.exists()

    def test_csv_de_entrada_tambem_funciona(self, cli_ctx: CLIContext, tmp_path: Path) -> None:
        entrada = tmp_path / "base.csv"
        entrada.write_text("codigo,uf,situacao,origem\n1,Sao Paulo,ATIVO,\n", encoding="utf-8")
        saida = tmp_path / "out"

        result = CliRunner().invoke(
            corrigir,
            [str(entrada), "--regras", str(REGRAS), "--out-dir", str(saida)],
            obj=cli_ctx,
        )
        assert result.exit_code == 0
        base = pd.read_csv(saida / CORRECTED_CSV_NAME, dtype=str).fillna("")
        assert base.loc[0, "uf"] == "SP"
        assert (saida / CORRECTED_XLSX_NAME).exists()


def test_arquivo_de_entrada_fica_intocado(cli_ctx: CLIContext, tmp_path: Path) -> None:
    antes = hashlib.sha256(BASE.read_bytes()).hexdigest()
    _run(cli_ctx, "--out-dir", str(tmp_path))
    assert hashlib.sha256(BASE.read_bytes()).hexdigest() == antes


def test_comando_registrado_na_cli() -> None:
    from autotarefas.cli.main import cli

    assert "corrigir" in cli.commands
