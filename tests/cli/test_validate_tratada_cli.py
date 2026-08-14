"""
Teste de ponta a ponta da planilha tratada (RF-PLA-009).

Um XLSX formatado entra pelo comando `validate --mode limpeza --out-dir` e
sai uma `planilha_tratada.xlsx` que o cliente reconhece como a dele — com
os valores normalizados e o relatorio do que nao pode ser preservado.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from click.testing import CliRunner
from openpyxl import load_workbook

from autotarefas.cli.commands.validate import PRESERVATION_REPORT_NAME, validate
from autotarefas.cli.context import CLIContext
from autotarefas.tasks.presentation import TREATED_XLSX_NAME

FIXTURES = Path(__file__).parent.parent / "fixtures" / "apresentacao"
ORIGINAL = FIXTURES / "original_formatado.xlsx"
SCHEMA = FIXTURES / "schema.yaml"


@pytest.fixture
def cli_ctx() -> CLIContext:
    return CLIContext()


def _run(out_dir: Path, ctx: CLIContext, *extra: str):
    return CliRunner().invoke(
        validate,
        [
            str(ORIGINAL),
            "--schema",
            str(SCHEMA),
            "--mode",
            "limpeza",
            "--out-dir",
            str(out_dir),
            *extra,
        ],
        obj=ctx,
    )


class TestPlanilhaTratada:
    def test_gera_a_planilha_tratada_e_o_relatorio(
        self, cli_ctx: CLIContext, tmp_path: Path
    ) -> None:
        result = _run(tmp_path, cli_ctx)
        assert result.exit_code == 0
        assert (tmp_path / TREATED_XLSX_NAME).exists()
        assert (tmp_path / PRESERVATION_REPORT_NAME).exists()

    def test_apresentacao_do_original_e_preservada(
        self, cli_ctx: CLIContext, tmp_path: Path
    ) -> None:
        _run(tmp_path, cli_ctx)
        ws = load_workbook(tmp_path / TREATED_XLSX_NAME)["Dados"]

        assert ws.freeze_panes == "A2"
        assert ws["A1"].fill.fgColor.rgb == "001F4E78"
        assert ws["D2"].number_format == "R$ #,##0.00"

    def test_valores_saem_normalizados(self, cli_ctx: CLIContext, tmp_path: Path) -> None:
        _run(tmp_path, cli_ctx)
        ws = load_workbook(tmp_path / TREATED_XLSX_NAME)["Dados"]

        assert ws["B2"].value == "Ana Lima"
        assert ws["C2"].value == "ana@x.com"

    def test_relatorio_de_preservacao(self, cli_ctx: CLIContext, tmp_path: Path) -> None:
        _run(tmp_path, cli_ctx)
        payload = json.loads((tmp_path / PRESERVATION_REPORT_NAME).read_text(encoding="utf-8"))

        assert payload["arquivo_original"] == "original_formatado.xlsx"
        assert payload["celulas_aplicadas"] >= 1
        assert payload["preservacao_total"] is True

    def test_original_intocado(self, cli_ctx: CLIContext, tmp_path: Path) -> None:
        antes = hashlib.sha256(ORIGINAL.read_bytes()).hexdigest()
        _run(tmp_path, cli_ctx)
        assert hashlib.sha256(ORIGINAL.read_bytes()).hexdigest() == antes

    def test_modo_auditoria_nao_gera_tratada(self, cli_ctx: CLIContext, tmp_path: Path) -> None:
        """Sem limpeza nao ha o que tratar: a tratada seria uma copia enganosa."""
        result = CliRunner().invoke(
            validate,
            [
                str(ORIGINAL),
                "--schema",
                str(SCHEMA),
                "--mode",
                "auditoria",
                "--out-dir",
                str(tmp_path),
            ],
            obj=cli_ctx,
        )
        assert result.exit_code == 0
        assert not (tmp_path / TREATED_XLSX_NAME).exists()

    def test_csv_nao_gera_tratada(self, cli_ctx: CLIContext, tmp_path: Path) -> None:
        csv = tmp_path / "dados.csv"
        csv.write_text("codigo,nome,email,valor\n1, Ana ,A@X.COM,10\n", encoding="utf-8")
        saida = tmp_path / "out"

        result = CliRunner().invoke(
            validate,
            [
                str(csv),
                "--schema",
                str(SCHEMA),
                "--mode",
                "limpeza",
                "--out-dir",
                str(saida),
            ],
            obj=cli_ctx,
        )
        assert result.exit_code == 0
        assert not (saida / TREATED_XLSX_NAME).exists()
