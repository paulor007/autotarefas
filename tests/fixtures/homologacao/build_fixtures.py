"""
Fixtures da homologacao do card de planilhas (jornada do Live).

A planilha real do proprietario estava limpa demais para provar as
correcoes: nenhuma celula precisou mudar. Estas fixtures sao pequenas,
sinteticas, deterministicas e SEM dados pessoais — cada linha existe para
provar um comportamento.

    A_vendas_com_anomalias.xlsx  correcoes seguras x itens de revisao
    B_duas_abas.xlsx             ambiguidade de aba (o sistema pergunta)
    C_apresentacao_rica.xlsx     preservacao (painel, filtro, formatos, formula)
    D_com_grafico.xlsx           limitacao declarada do openpyxl

Rodar:  python tests/fixtures/homologacao/build_fixtures.py
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

AQUI = Path(__file__).parent

CABECALHO = [
    "Codigo Venda",
    "Data",
    "Loja",
    "Produto",
    "Quantidade",
    "Valor Unitario",
    "Valor Final",
]

#: Cada linha carrega UM caso, documentado no comentario ao lado.
#: A coluna "Codigo Venda" repete de proposito em 4 linhas (uma venda com
#: varios itens): isso NAO pode virar duplicidade.
LINHAS_A: list[list[object]] = [
    CABECALHO,
    # 2: limpa (controle)
    ["V-001", "01/12/2026", "Loja Centro", "Caneta", 2, "10,00", "20,00"],
    # 3: espaco no inicio e no fim + espacos internos duplicados
    ["V-002", "02/12/2026", "  Loja   Norte  ", " Caderno  ", 1, "15,00", "15,00"],
    # 4: codigo com zeros a esquerda, guardado como TEXTO
    ["00123", "03/12/2026", "Loja Sul", "Borracha", 3, "2,50", "7,50"],
    # 5: quantidade guardada como texto
    ["V-004", "04/12/2026", "Loja Centro", "Regua", "4", "5,00", "20,00"],
    # 6: data invalida (nao existe 32/12)
    ["V-005", "32/12/2026", "Loja Centro", "Estojo", 1, "30,00", "30,00"],
    # 7: quantidade zero
    ["V-006", "05/12/2026", "Loja Norte", "Mochila", 0, "120,00", "0,00"],
    # 8: quantidade negativa
    ["V-007", "06/12/2026", "Loja Sul", "Lapis", -2, "3,00", "-6,00"],
    # 9: valor unitario negativo
    ["V-008", "07/12/2026", "Loja Centro", "Apontador", 1, "-4,00", "-4,00"],
    # 10: Valor Final != Quantidade x Valor Unitario (2 x 10 = 20, nao 25)
    ["V-009", "08/12/2026", "Loja Norte", "Tesoura", 2, "10,00", "25,00"],
    # 11 e 12: LINHA COMPLETAMENTE DUPLICADA (o par de repetidas)
    ["V-010", "09/12/2026", "Loja Sul", "Cola", 1, "8,00", "8,00"],
    ["V-010", "09/12/2026", "Loja Sul", "Cola", 1, "8,00", "8,00"],
    # 13: campo obrigatorio vazio (Produto)
    ["V-011", "10/12/2026", "Loja Centro", "", 1, "9,00", "9,00"],
    # 14: ambiguo — "1.234" pode ser mil-duzentos-e-trinta-e-quatro ou 1,234.
    # O AutoTarefas NAO decide: fica como esta e vai para revisao humana.
    ["V-012", "11/12/2026", "Loja Norte", "Pasta", 1, "1.234", "1.234"],
    # 15-18: MESMO codigo de venda em 4 itens diferentes. Chave repetida
    # esperada — nao e duplicidade e nao pode virar falso positivo.
    ["V-100", "12/12/2026", "Loja Centro", "Item A", 1, "10,00", "10,00"],
    ["V-100", "12/12/2026", "Loja Centro", "Item B", 2, "20,00", "40,00"],
    ["V-100", "12/12/2026", "Loja Centro", "Item C", 3, "30,00", "90,00"],
    ["V-100", "12/12/2026", "Loja Centro", "Item D", 4, "40,00", "160,00"],
]

LINHAS_B1 = [
    ["Codigo", "Produto", "Quantidade"],
    ["A-1", "Caneta", 10],
    ["A-2", "Caderno", 20],
    ["A-3", "Borracha", 30],
]

LINHAS_B2 = [
    ["Codigo", "Produto", "Quantidade"],
    ["B-1", "Mochila", 40],
    ["B-2", "Estojo", 50],
    ["B-3", "Regua", 60],
]


def _fixture_a() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    for linha in LINHAS_A:
        ws.append(linha)
    # A coluna do codigo e TEXTO: "00123" nao pode virar 123.
    for linha in range(2, ws.max_row + 1):
        ws.cell(row=linha, column=1).number_format = "@"
    wb.save(AQUI / "A_vendas_com_anomalias.xlsx")


def _fixture_b() -> None:
    """Duas abas igualmente plausiveis: o sistema tem de perguntar."""
    wb = Workbook()
    primeira = wb.active
    primeira.title = "Dezembro"
    for linha in LINHAS_B1:
        primeira.append(linha)
    segunda = wb.create_sheet("Janeiro")
    for linha in LINHAS_B2:
        segunda.append(linha)
    wb.save(AQUI / "B_duas_abas.xlsx")


def _fixture_c() -> None:
    """Tudo o que a preservacao precisa provar, num arquivo so."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Base"
    ws.append(["Codigo", "Cliente", "Data", "Quantidade", "Valor", "Total"])
    ws.append(["00123", "  Cliente   Um  ", datetime(2026, 12, 1), 2, 10.5, "=D2*E2"])
    ws.append(["00124", "Cliente Dois", datetime(2026, 12, 2), 3, 20.0, "=D3*E3"])

    cabecalho_fill = PatternFill("solid", fgColor="1F4E78")
    cabecalho_font = Font(name="Arial", bold=True, color="FFFFFF")
    for celula in ws[1]:
        celula.font = cabecalho_font
        celula.fill = cabecalho_fill
        celula.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{ws.max_row}"
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 14
    for linha in range(2, ws.max_row + 1):
        ws.cell(row=linha, column=1).number_format = "@"
        ws.cell(row=linha, column=3).number_format = "DD/MM/YYYY"
        ws.cell(row=linha, column=5).number_format = "R$ #,##0.00"
        ws.cell(row=linha, column=6).number_format = "R$ #,##0.00"
    wb.save(AQUI / "C_apresentacao_rica.xlsx")


def _fixture_d() -> None:
    """Grafico: o openpyxl nao reescreve, e o relatorio precisa declarar."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Base"
    ws.append(["Produto", "Quantidade"])
    ws.append(["Caneta", 10])
    ws.append(["Caderno", 20])
    grafico = BarChart()
    grafico.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    ws.add_chart(grafico, "D2")
    wb.save(AQUI / "D_com_grafico.xlsx")


def main() -> None:
    _fixture_a()
    _fixture_b()
    _fixture_c()
    _fixture_d()
    print(f"OK: 4 fixtures de homologacao em {AQUI}")


if __name__ == "__main__":
    main()
