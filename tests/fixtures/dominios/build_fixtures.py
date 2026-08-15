"""
Fixtures de DOMÍNIOS diferentes para o Card 01.

A planilha de vendas do proprietário é um teste real — não pode virar a
base do funcionamento. Estas fixtures são pequenas, sintéticas,
determinísticas e sem dados pessoais, e cobrem contextos distintos para
provar que o card não foi calibrado para um único formato.

    vendas_simples.xlsx        XLSX sem formatação nenhuma (deve sugerir organização)
    financeiro_profissional.xlsx  já organizado (NÃO pode ser reformatado)
    servico_publico.xlsx       protocolos, órgãos, datas — domínio público
    estoque_com_formulas.xlsx  fórmulas (ordenação deve ser recusada)
    clientes.csv               CSV genérico
    pesquisa_ambigua.xlsx      capa + mesclagens: estrutura ambígua
    atendimentos_duas_abas.xlsx  duas abas tabulares candidatas
    contratos_sem_anomalia.xlsx  nada a corrigir e nada a organizar

Rodar:  python tests/fixtures/dominios/build_fixtures.py
"""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

AQUI = Path(__file__).parent

_CABECALHO_FILL = PatternFill("solid", fgColor="1F4E78")
_CABECALHO_FONT = Font(bold=True, color="FFFFFF")


def _profissionalizar(ws: Worksheet, colunas: int) -> None:
    """Deixa a aba com a cara que o card considera 'já organizada'."""
    for celula in ws[1]:
        celula.font = _CABECALHO_FONT
        celula.fill = _CABECALHO_FILL
        celula.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(colunas)}{ws.max_row}"
    for indice in range(1, colunas + 1):
        letra = get_column_letter(indice)
        maior = max(
            (len(str(ws.cell(row=r, column=indice).value or "")) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[letra].width = max(12, min(maior + 2, 40))


def _vendas_simples() -> None:
    """Dados bons, apresentação crua: o caso clássico de 'pode melhorar'."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    ws.append(["Pedido", "Data", "Vendedor", "Produto", "Quantidade", "Valor"])
    base = [
        ("P-001", datetime(2026, 1, 5), "Ana", "Teclado", 2, 150.0),
        ("P-002", datetime(2026, 1, 7), "Bruno", "Monitor", 1, 900.0),
        ("P-003", datetime(2026, 1, 9), "Ana", "Mouse", 5, 45.0),
        ("P-004", datetime(2026, 2, 2), "Carla", "Teclado", 3, 150.0),
        ("P-005", datetime(2026, 2, 14), "Bruno", "Headset", 2, 220.0),
        ("P-006", datetime(2026, 2, 20), "Ana", "Monitor", 1, 900.0),
        ("P-007", datetime(2026, 3, 3), "Carla", "Mouse", 4, 45.0),
        ("P-008", datetime(2026, 3, 11), "Bruno", "Teclado", 1, 150.0),
        ("P-009", datetime(2026, 3, 18), "Ana", "Headset", 2, 220.0),
        ("P-010", datetime(2026, 3, 25), "Carla", "Monitor", 2, 900.0),
    ]
    # 20 linhas: acima do limiar de filtro/painel, para que a avaliacao cobre.
    for repeticao in range(2):
        for pedido, data, vendedor, produto, quantidade, valor in base:
            ws.append(
                [
                    f"{pedido}-{repeticao}",
                    data,
                    vendedor,
                    produto,
                    quantidade,
                    valor,
                ]
            )
    wb.save(AQUI / "vendas_simples.xlsx")


def _financeiro_profissional() -> None:
    """Já organizado: o card não pode propor reformatação."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Lancamentos"
    ws.append(["Documento", "Competencia", "Centro de Custo", "Natureza", "Valor"])
    linhas = [
        ("DOC-100", datetime(2026, 1, 31), "Administrativo", "Despesa", 1200.55),
        ("DOC-101", datetime(2026, 1, 31), "Comercial", "Receita", 8400.00),
        ("DOC-102", datetime(2026, 2, 28), "Administrativo", "Despesa", 990.10),
        ("DOC-103", datetime(2026, 2, 28), "Operacoes", "Despesa", 2310.75),
        ("DOC-104", datetime(2026, 3, 31), "Comercial", "Receita", 9100.00),
        ("DOC-105", datetime(2026, 3, 31), "Operacoes", "Despesa", 1750.40),
    ]
    for _ in range(3):
        for linha in linhas:
            ws.append(list(linha))
    for linha in range(2, ws.max_row + 1):
        ws.cell(row=linha, column=2).number_format = "DD/MM/YYYY"
        ws.cell(row=linha, column=5).number_format = "R$ #,##0.00"
        ws.cell(row=linha, column=5).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=linha, column=2).alignment = Alignment(horizontal="center", vertical="center")
        for coluna in (1, 3, 4):
            ws.cell(row=linha, column=coluna).alignment = Alignment(
                horizontal="left", vertical="center"
            )
    _profissionalizar(ws, 5)
    wb.save(AQUI / "financeiro_profissional.xlsx")


def _servico_publico() -> None:
    """Protocolos com zeros à esquerda, órgão e situação."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Protocolos"
    ws.append(["Protocolo", "Abertura", "Orgao", "Assunto", "Situacao", "Dias"])
    linhas = [
        ("000123", datetime(2026, 1, 8), "Secretaria de Saude", "Agendamento", "Concluido", 4),
        ("000124", datetime(2026, 1, 9), "Secretaria de Obras", "Iluminacao", "Em analise", 12),
        ("000125", datetime(2026, 1, 10), "Secretaria de Saude", "Medicamento", "Concluido", 2),
        ("000126", datetime(2026, 1, 15), "Ouvidoria", "Reclamacao", "Pendente", 30),
        ("000127", datetime(2026, 2, 1), "Secretaria de Obras", "Buraco na via", "Em analise", 18),
        ("000128", datetime(2026, 2, 3), "Ouvidoria", "Elogio", "Concluido", 1),
        ("000129", datetime(2026, 2, 11), "Secretaria de Saude", "Agendamento", "Pendente", 9),
        ("000130", datetime(2026, 2, 19), "Secretaria de Obras", "Iluminacao", "Concluido", 6),
        # Linha 10 e 11: repetida de proposito (duplicidade sempre verificada).
        ("000131", datetime(2026, 3, 2), "Ouvidoria", "Reclamacao", "Pendente", 21),
        ("000131", datetime(2026, 3, 2), "Ouvidoria", "Reclamacao", "Pendente", 21),
        ("000132", datetime(2026, 3, 8), "Secretaria de Saude", "  Exame   ", "Concluido", 3),
    ]
    for linha in linhas:
        ws.append(list(linha))
    for linha in range(2, ws.max_row + 1):
        ws.cell(row=linha, column=1).number_format = "@"
    wb.save(AQUI / "servico_publico.xlsx")


def _estoque_com_formulas() -> None:
    """Fórmulas: a ordenação tem de ser recusada, não aplicada."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Estoque"
    ws.append(["SKU", "Descricao", "Quantidade", "Custo Unitario", "Total"])
    dados = [
        ("SKU-9", "Parafuso", 120, 0.35),
        ("SKU-3", "Porca", 340, 0.20),
        ("SKU-7", "Arruela", 80, 0.15),
        ("SKU-1", "Prego", 500, 0.05),
    ]
    for indice, (sku, descricao, quantidade, custo) in enumerate(dados, start=2):
        ws.append([sku, descricao, quantidade, custo, f"=C{indice}*D{indice}"])
    wb.save(AQUI / "estoque_com_formulas.xlsx")


def _clientes_csv() -> None:
    """CSV genérico: sem apresentação para avaliar."""
    linhas = [
        ["Codigo", "Nome", "Cidade", "Cadastro", "Ativo"],
        ["C-001", "Cliente Alfa", "Recife", "2026-01-10", "sim"],
        ["C-002", "Cliente Beta", "Curitiba", "2026-01-14", "sim"],
        ["C-003", "  Cliente   Gama  ", "Sao Paulo", "2026-02-01", "nao"],
        ["C-004", "Cliente Delta", "Recife", "2026-02-09", "sim"],
        ["C-005", "Cliente Epsilon", "Belem", "2026-03-05", "sim"],
    ]
    with (AQUI / "clientes.csv").open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle, lineterminator="\n").writerows(linhas)


def _pesquisa_ambigua() -> None:
    """Capa mesclada por cima da tabela: estrutura ambígua."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Pesquisa"
    ws["A1"] = "PESQUISA DE SATISFACAO - 1o TRIMESTRE"
    ws.merge_cells("A1:D1")
    ws.append([])
    ws.append(["Respondente", "Nota", "Comentario", "Data"])
    ws.append([1, 9, "Bom atendimento", datetime(2026, 1, 5)])
    ws.append([2, 7, "", datetime(2026, 1, 6)])
    ws.append([3, 10, "Excelente", datetime(2026, 1, 7)])
    wb.save(AQUI / "pesquisa_ambigua.xlsx")


def _atendimentos_duas_abas() -> None:
    """Duas abas tabulares: o card precisa perguntar qual processar."""
    wb = Workbook()
    janeiro = wb.active
    janeiro.title = "Janeiro"
    janeiro.append(["Atendimento", "Canal", "Duracao", "Resolvido"])
    for indice in range(1, 13):
        janeiro.append([f"A-{indice:03d}", "Telefone", 5 + indice, "sim"])

    fevereiro = wb.create_sheet("Fevereiro")
    fevereiro.append(["Atendimento", "Canal", "Duracao", "Resolvido"])
    for indice in range(1, 13):
        fevereiro.append([f"B-{indice:03d}", "Chat", 3 + indice, "nao"])

    capa = wb.create_sheet("Leia-me")
    capa["A1"] = "Relatorio de atendimentos"
    capa["A3"] = "Gerado pelo sistema interno"

    wb.create_sheet("Rascunho")  # aba vazia
    wb.save(AQUI / "atendimentos_duas_abas.xlsx")


def _contratos_sem_anomalia() -> None:
    """Dados limpos e apresentação boa: o card não deve propor nada."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Contratos"
    ws.append(["Contrato", "Fornecedor", "Inicio", "Valor Mensal"])
    linhas = [
        ("CT-001", "Fornecedor A", datetime(2026, 1, 1), 1500.0),
        ("CT-002", "Fornecedor B", datetime(2026, 2, 1), 2300.0),
        ("CT-003", "Fornecedor C", datetime(2026, 3, 1), 980.0),
    ]
    for linha in linhas:
        ws.append(list(linha))
    for linha in range(2, ws.max_row + 1):
        ws.cell(row=linha, column=3).number_format = "DD/MM/YYYY"
        ws.cell(row=linha, column=3).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=linha, column=4).number_format = "R$ #,##0.00"
        ws.cell(row=linha, column=4).alignment = Alignment(horizontal="right", vertical="center")
        for coluna in (1, 2):
            ws.cell(row=linha, column=coluna).alignment = Alignment(
                horizontal="left", vertical="center"
            )
    _profissionalizar(ws, 4)
    wb.save(AQUI / "contratos_sem_anomalia.xlsx")


def main() -> None:
    _vendas_simples()
    _financeiro_profissional()
    _servico_publico()
    _estoque_com_formulas()
    _clientes_csv()
    _pesquisa_ambigua()
    _atendimentos_duas_abas()
    _contratos_sem_anomalia()
    print(f"OK: 8 fixtures de dominios em {AQUI}")


if __name__ == "__main__":
    main()
