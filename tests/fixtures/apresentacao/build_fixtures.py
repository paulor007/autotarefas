"""
Gerador da fixture de APRESENTACAO (RF-PLA-009).

`original_formatado.xlsx` tem tudo o que o cliente reconhece como "a
planilha dele": cabecalho colorido, painel congelado, autofiltro, largura
de coluna ajustada e formato de moeda — e, de proposito, valores sujos
(espacos extras, moeda como texto) para que a limpeza tenha o que tratar.

A versao tratada gerada pelo AutoTarefas precisa sair com a MESMA cara.

Rodar:  python tests/fixtures/apresentacao/build_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

AQUI = Path(__file__).parent

CABECALHO_FILL = PatternFill("solid", fgColor="1F4E78")
CABECALHO_FONT = Font(name="Arial", bold=True, color="FFFFFF")

LINHAS: list[list[str]] = [
    ["codigo", "nome", "email", "valor"],
    ["00123", "  Ana   Lima  ", " ANA@X.COM ", "1.234,56"],
    ["00124", "Bruno Sa", "bruno@x.com", "10,00"],
    ["00125", "  Carla Nunes", "carla@x.com", "99,90"],
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    for linha in LINHAS:
        ws.append(linha)

    for celula in ws[1]:
        celula.font = CABECALHO_FONT
        celula.fill = CABECALHO_FILL
        celula.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{ws.max_row}"
    ws.column_dimensions["B"].width = 40
    for linha in range(2, ws.max_row + 1):
        ws.cell(row=linha, column=4).number_format = "R$ #,##0.00"

    destino = AQUI / "original_formatado.xlsx"
    wb.save(destino)
    print(f"OK: fixture de apresentacao em {destino}")


if __name__ == "__main__":
    main()
