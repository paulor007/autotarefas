"""
Testes da preservacao da apresentacao original (RF-PLA-009).

Criterio de aceite da ficha: cabecalho colorido, painel congelado e
formato de moeda do original presentes na tratada — e o significado dos
dados jamais alterado pela formatacao.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

from autotarefas.tasks.presentation import (
    TREATED_XLSX_NAME,
    supports_presentation,
    write_treated_xlsx,
)

FIXTURE = Path(__file__).parent.parent / "fixtures" / "apresentacao" / "original_formatado.xlsx"


@pytest.fixture
def mudancas() -> list[dict[str, object]]:
    return [
        {"line": 2, "column": "nome", "before": "  Ana   Lima  ", "after": "Ana Lima"},
        {"line": 2, "column": "email", "before": " ANA@X.COM ", "after": "ana@x.com"},
    ]


class TestPreservacao:
    def test_apresentacao_do_original_sobrevive(
        self, mudancas: list[dict[str, object]], tmp_path: Path
    ) -> None:
        destino = tmp_path / TREATED_XLSX_NAME
        write_treated_xlsx(FIXTURE, destino, mudancas, sheet="Dados")

        ws = load_workbook(destino)["Dados"]
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref == "A1:D4"
        assert ws.column_dimensions["B"].width == 40
        assert ws["A1"].fill.fgColor.rgb == "001F4E78"
        assert ws["A1"].font.bold is True
        assert ws["D2"].number_format == "R$ #,##0.00"

    def test_valores_tratados_sao_aplicados(
        self, mudancas: list[dict[str, object]], tmp_path: Path
    ) -> None:
        destino = tmp_path / TREATED_XLSX_NAME
        relatorio = write_treated_xlsx(FIXTURE, destino, mudancas, sheet="Dados")

        ws = load_workbook(destino)["Dados"]
        assert ws["B2"].value == "Ana Lima"
        assert ws["C2"].value == "ana@x.com"
        assert relatorio.applied == 2

    def test_celulas_nao_tratadas_ficam_iguais(
        self, mudancas: list[dict[str, object]], tmp_path: Path
    ) -> None:
        destino = tmp_path / TREATED_XLSX_NAME
        write_treated_xlsx(FIXTURE, destino, mudancas, sheet="Dados")

        ws = load_workbook(destino)["Dados"]
        assert ws["B3"].value == "Bruno Sa"
        assert ws["B4"].value == "  Carla Nunes"  # nao estava na lista de mudancas

    def test_original_nunca_e_alterado(
        self, mudancas: list[dict[str, object]], tmp_path: Path
    ) -> None:
        antes = hashlib.sha256(FIXTURE.read_bytes()).hexdigest()
        write_treated_xlsx(FIXTURE, tmp_path / TREATED_XLSX_NAME, mudancas, sheet="Dados")
        assert hashlib.sha256(FIXTURE.read_bytes()).hexdigest() == antes

    def test_valor_tipado_vem_do_dataframe(self, tmp_path: Path) -> None:
        """O audit trail guarda texto; o Excel precisa do numero."""
        frame = pd.DataFrame(
            {"codigo": ["00123"], "nome": ["Ana Lima"], "email": ["a@x.com"], "valor": [1234.56]}
        )
        destino = tmp_path / TREATED_XLSX_NAME
        write_treated_xlsx(
            FIXTURE,
            destino,
            [{"line": 2, "column": "valor", "after": "1234.56"}],
            dataframe=frame,
            sheet="Dados",
        )

        valor = load_workbook(destino)["Dados"]["D2"].value
        assert valor == 1234.56
        assert isinstance(valor, float)


class TestLimitesDeclarados:
    def test_grafico_do_original_e_declarado_como_perdido(self, tmp_path: Path) -> None:
        origem = tmp_path / "com_grafico.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Dados"
        ws.append(["codigo", "valor"])
        ws.append(["1", 10])
        ws.append(["2", 20])
        grafico = BarChart()
        grafico.add_data(Reference(ws, min_col=2, min_row=1, max_row=3))
        ws.add_chart(grafico, "D2")
        wb.save(origem)

        relatorio = write_treated_xlsx(origem, tmp_path / "tratada.xlsx", [], sheet="Dados")

        assert relatorio.fully_preserved is False
        assert "grafico" in relatorio.not_preserved[0]

    def test_mudanca_em_coluna_inexistente_e_reportada(self, tmp_path: Path) -> None:
        destino = tmp_path / TREATED_XLSX_NAME
        relatorio = write_treated_xlsx(
            FIXTURE,
            destino,
            [{"line": 2, "column": "fantasma", "after": "x"}],
            sheet="Dados",
        )
        assert relatorio.applied == 0
        assert "fantasma" in relatorio.skipped[0]

    def test_csv_nao_tem_apresentacao_para_preservar(self, tmp_path: Path) -> None:
        csv = tmp_path / "dados.csv"
        csv.write_text("a,b\n1,2\n", encoding="utf-8")

        assert supports_presentation(csv) is False
        with pytest.raises(ValueError, match="xlsx"):
            write_treated_xlsx(csv, tmp_path / "saida.xlsx", [])

    def test_relatorio_serializa_para_json(
        self, mudancas: list[dict[str, object]], tmp_path: Path
    ) -> None:
        relatorio = write_treated_xlsx(
            FIXTURE, tmp_path / TREATED_XLSX_NAME, mudancas, sheet="Dados"
        )
        payload = relatorio.as_dict()
        assert payload["celulas_aplicadas"] == 2
        assert payload["preservacao_total"] is True
        assert payload["arquivo_original"] == "original_formatado.xlsx"
