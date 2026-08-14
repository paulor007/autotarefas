"""
Testes da apresentacao comum dos artefatos XLSX.

Este modulo e compartilhado por `planilha_validada.xlsx` (PLA-007) e
`divergencias.xlsx` (REC-001): um defeito aqui aparece nos dois artefatos
ao mesmo tempo, entao os casos de borda ficam cobertos aqui, uma vez so.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

from autotarefas.tasks.xlsx_style import (
    autofit_columns,
    cell_value,
    style_header_row,
    write_dataframe_sheet,
)


class TestCellValue:
    def test_none_e_nan_viram_texto_vazio(self) -> None:
        assert cell_value(None) == ""
        assert cell_value(float("nan")) == ""

    def test_tipos_nativos_passam_direto(self) -> None:
        assert cell_value("00123") == "00123"
        assert cell_value(42) == 42
        assert cell_value(True) is True

    def test_numpy_vira_tipo_nativo(self) -> None:
        convertido = cell_value(np.int64(7))
        assert convertido == 7
        assert isinstance(convertido, int)

    def test_objeto_desconhecido_vira_texto(self) -> None:
        class Estranho:
            def __str__(self) -> str:
                return "estranho"

        assert cell_value(Estranho()) == "estranho"


class TestFormatacao:
    def test_cabecalho_congela_painel_e_liga_filtro(self, tmp_path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["a", "b"])
        ws.append(["1", "2"])

        style_header_row(ws, 2)

        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref == "A1:B2"

    def test_planilha_sem_colunas_nao_quebra(self) -> None:
        wb = Workbook()
        ws = wb.active
        style_header_row(ws, 0)
        assert ws.auto_filter.ref is None

    def test_largura_respeita_o_teto(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["x" * 200])
        autofit_columns(ws, max_width=30)
        assert ws.column_dimensions["A"].width == 30

    def test_dataframe_vira_aba_formatada(self, tmp_path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        write_dataframe_sheet(ws, pd.DataFrame({"nome": ["Ana"], "valor": [None]}))

        destino = tmp_path / "saida.xlsx"
        wb.save(destino)

        lido = load_workbook(destino).active
        assert [c.value for c in lido[1]] == ["nome", "valor"]
        assert lido["B2"].value is None  # celula vazia gravada como ""
