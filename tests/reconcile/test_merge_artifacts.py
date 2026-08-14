"""Testes dos artefatos da reconciliacao (RF-REC-002)."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from autotarefas.reconcile.compare import compare_tables
from autotarefas.reconcile.merge import ReconcilePolicy, ReconciliationResult, reconcile_tables
from autotarefas.reconcile.merge_artifacts import (
    BASE_CSV_NAME,
    BASE_XLSX_NAME,
    JSON_REPORT_NAME,
    REVIEW_XLSX_NAME,
    generate_summary,
    write_reconciliation_artifacts,
)
from autotarefas.reconcile.result import TableSource
from autotarefas.reconcile.tolerance import Tolerance

COLUNAS = ["codigo", "nome", "valor"]


def _tabela(nome: str, linhas: list[list[str]]) -> TableSource:
    return TableSource(name=nome, frame=pd.DataFrame(linhas, columns=COLUNAS, dtype=str))


@pytest.fixture
def resultado() -> ReconciliationResult:
    a = _tabela(
        "a.xlsx",
        [
            ["1", "Ana", "1.000,00"],
            ["2", "Bruno", "20,00"],
            ["3", "Carla", "30,00"],
            ["5", "Eva", "50,00"],
            ["5", "Eva II", "51,00"],
        ],
    )
    b = _tabela(
        "b.xlsx",
        [
            ["1", "Ana", "1.000,01"],
            ["2", "Bruno Silva", "25,00"],
            ["4", "Diego", "40,00"],
        ],
    )
    comparacao = compare_tables(
        a, b, key_columns=["codigo"], tolerances=[Tolerance("valor", "absoluta", 0.01)]
    )
    return reconcile_tables(a, b, comparacao, ReconcilePolicy(updatable_columns=("valor",)))


class TestArtefatos:
    def test_gera_os_quatro_arquivos(self, resultado: ReconciliationResult, tmp_path: Path) -> None:
        caminhos = write_reconciliation_artifacts(resultado, tmp_path)
        assert [p.name for p in caminhos] == [
            BASE_XLSX_NAME,
            BASE_CSV_NAME,
            REVIEW_XLSX_NAME,
            JSON_REPORT_NAME,
        ]
        assert all(p.exists() for p in caminhos)

    def test_base_csv_tem_os_valores_decididos(
        self, resultado: ReconciliationResult, tmp_path: Path
    ) -> None:
        write_reconciliation_artifacts(resultado, tmp_path)
        base = pd.read_csv(tmp_path / BASE_CSV_NAME, dtype=str).fillna("")

        assert list(base.columns) == [*COLUNAS, "_origem"]
        assert base.loc[base["codigo"] == "2", "valor"].item() == "25,00"
        assert base.loc[base["codigo"] == "2", "nome"].item() == "Bruno"
        assert "5" not in list(base["codigo"])  # chave duplicada nao entra

    def test_base_xlsx_tem_abas_de_base_decisoes_e_resumo(
        self, resultado: ReconciliationResult, tmp_path: Path
    ) -> None:
        write_reconciliation_artifacts(resultado, tmp_path)
        wb = load_workbook(tmp_path / BASE_XLSX_NAME)
        assert wb.sheetnames == ["Base conciliada", "Decisoes", "Resumo"]

        decisoes = wb["Decisoes"]
        cabecalho = [c.value for c in decisoes[1]]
        assert cabecalho == [
            "codigo",
            "coluna",
            "valor escolhido",
            "valor descartado",
            "fonte",
            "regra",
        ]

    def test_revisao_lista_o_conflito_real(
        self, resultado: ReconciliationResult, tmp_path: Path
    ) -> None:
        write_reconciliation_artifacts(resultado, tmp_path)
        ws = load_workbook(tmp_path / REVIEW_XLSX_NAME)["Para revisao"]
        linhas = [tuple(linha) for linha in ws.iter_rows(min_row=2, values_only=True)]

        motivos = {linha[1] for linha in linhas}
        assert "divergencia em campo nao autorizado" in motivos
        assert "chave repetida na fonte (nao pareada)" in motivos

    def test_relatorio_json_registra_politica_e_decisoes(
        self, resultado: ReconciliationResult, tmp_path: Path
    ) -> None:
        write_reconciliation_artifacts(resultado, tmp_path, {"task_name": "conciliar"})
        payload = json.loads((tmp_path / JSON_REPORT_NAME).read_text(encoding="utf-8"))

        assert payload["task_name"] == "conciliar"
        assert payload["requisito"] == "RF-REC-002"
        assert payload["politica"]["fonte_principal"] == "a"
        assert payload["politica"]["campos_atualizaveis"] == ["valor"]
        assert payload["resumo"]["campos_atualizados"] == 1
        assert payload["resumo"]["toleradas"] == 1
        assert any(d["regra"] == "tolerancia" for d in payload["decisoes"])
        assert any(i["motivo"] == "chave_duplicada" for i in payload["para_revisao"])
        assert payload["limitacoes"]

    def test_tolerancia_aparece_no_relatorio(
        self, resultado: ReconciliationResult, tmp_path: Path
    ) -> None:
        write_reconciliation_artifacts(resultado, tmp_path)
        payload = json.loads((tmp_path / JSON_REPORT_NAME).read_text(encoding="utf-8"))
        assert payload["tolerancias"] == ["valor: ate 0.01 (absoluta)"]


class TestResumoTextual:
    def test_mostra_decisoes_e_pendencias(self, resultado: ReconciliationResult) -> None:
        texto = generate_summary(resultado)

        assert "Reconciliacao por chave: codigo" in texto
        assert "fonte principal: A" in texto
        assert "'20,00'  ->  '25,00'" in texto
        assert "Para revisao" in texto
        assert "Limitacoes desta reconciliacao" in texto

    def test_trunca_listas_longas(self) -> None:
        a = _tabela("a.xlsx", [[str(n), f"Nome {n}", "1,00"] for n in range(6)])
        b = _tabela("b.xlsx", [[str(n), f"Nome {n}", "2,00"] for n in range(6)])
        comparacao = compare_tables(a, b, key_columns=["codigo"])
        resultado = reconcile_tables(
            a, b, comparacao, ReconcilePolicy(updatable_columns=("valor",))
        )

        texto = generate_summary(resultado, max_rows=2)
        assert "... e mais 4" in texto
