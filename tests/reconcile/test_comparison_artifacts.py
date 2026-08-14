"""
Testes dos artefatos da comparacao (RF-REC-001).

O que se cobra aqui e o que o operador recebe: os quatro arquivos com os
nomes da ficha, com as linhas inteiras das fontes, os valores ORIGINAIS e
as limitacoes declaradas dentro do proprio relatorio.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from autotarefas.reconcile.artifacts import (
    JSON_REPORT_NAME,
    ONLY_A_CSV_NAME,
    ONLY_B_CSV_NAME,
    ROW_COLUMN,
    XLSX_NAME,
    build_report_payload,
    generate_summary,
    write_comparison_artifacts,
)
from autotarefas.reconcile.compare import compare_tables
from autotarefas.reconcile.result import ComparisonResult, TableSource

COLUNAS = ["codigo", "nome", "valor"]


def _tabela(nome: str, linhas: list[list[str]], colunas: list[str] | None = None) -> TableSource:
    return TableSource(
        name=nome,
        frame=pd.DataFrame(linhas, columns=colunas or COLUNAS, dtype=str),
    )


@pytest.fixture
def tabelas() -> tuple[TableSource, TableSource]:
    a = _tabela(
        "a.xlsx",
        [
            ["1", "Ana", "10,00"],
            ["2", "Bruno", "20,00"],
            ["3", "Carla", "30,00"],
            ["4", "Diego", "40,00"],
            ["4", "Diego II", "41,00"],
        ],
    )
    b = _tabela(
        "b.xlsx",
        [
            ["1", "Ana", "10,00"],
            ["2", "Bruno", "25,00"],
            ["5", "Elis", "50,00"],
            ["4", "Diego", "40,00"],
        ],
    )
    return a, b


@pytest.fixture
def resultado(tabelas: tuple[TableSource, TableSource]) -> ComparisonResult:
    a, b = tabelas
    return compare_tables(a, b, key_columns=["codigo"])


class TestPayloadJSON:
    def test_estrutura_do_relatorio(self, resultado: ComparisonResult) -> None:
        payload = build_report_payload(resultado)

        assert payload["requisito"] == "RF-REC-001"
        assert payload["chave"] == ["codigo"]
        assert payload["colunas_comparadas"] == ["nome", "valor"]
        assert payload["resumo"]["divergente"] == 1
        assert payload["fonte_a"]["arquivo"] == "a.xlsx"
        assert payload["limitacoes"]

    def test_divergencia_traz_chave_linhas_e_valores(self, resultado: ComparisonResult) -> None:
        (divergencia,) = resultado_divergencias(resultado)
        assert divergencia["chave"] == {"codigo": "2"}
        assert divergencia["linha_a"] == 3
        assert divergencia["diferencas"] == [{"coluna": "valor", "a": "20,00", "b": "25,00"}]

    def test_conflito_lista_as_linhas_dos_dois_lados(self, resultado: ComparisonResult) -> None:
        (conflito,) = build_report_payload(resultado)["conflitos"]
        assert conflito["motivo"] == "chave_duplicada"
        assert conflito["linhas_a"] == [5, 6]
        assert conflito["linhas_b"] == [5]

    def test_meta_entra_no_topo(self, resultado: ComparisonResult) -> None:
        payload = build_report_payload(resultado, {"task_name": "comparar", "status": "success"})
        assert payload["task_name"] == "comparar"
        assert payload["status"] == "success"

    def test_identicos_nao_sao_listados(self, resultado: ComparisonResult) -> None:
        """Contador sim, lista nao: o arquivo e para decidir, nao para arquivar."""
        payload = build_report_payload(resultado)
        assert payload["resumo"]["identico"] == 1
        assert "identicos" not in payload


def resultado_divergencias(result: ComparisonResult) -> list[dict[str, object]]:
    divergencias: list[dict[str, object]] = build_report_payload(result)["divergencias"]
    return divergencias


class TestArtefatosEmDisco:
    def test_gera_os_quatro_arquivos_com_os_nomes_da_ficha(
        self,
        resultado: ComparisonResult,
        tabelas: tuple[TableSource, TableSource],
        tmp_path: Path,
    ) -> None:
        a, b = tabelas
        caminhos = write_comparison_artifacts(resultado, a, b, tmp_path)

        assert [p.name for p in caminhos] == [
            JSON_REPORT_NAME,
            XLSX_NAME,
            ONLY_A_CSV_NAME,
            ONLY_B_CSV_NAME,
        ]
        assert all(p.exists() for p in caminhos)

    def test_json_e_lido_de_volta(
        self,
        resultado: ComparisonResult,
        tabelas: tuple[TableSource, TableSource],
        tmp_path: Path,
    ) -> None:
        a, b = tabelas
        write_comparison_artifacts(resultado, a, b, tmp_path)

        payload = json.loads((tmp_path / JSON_REPORT_NAME).read_text(encoding="utf-8"))
        assert payload["resumo"]["somente_a"] == 1
        assert payload["resumo"]["somente_b"] == 1

    def test_csv_traz_a_linha_inteira_e_a_origem(
        self,
        resultado: ComparisonResult,
        tabelas: tuple[TableSource, TableSource],
        tmp_path: Path,
    ) -> None:
        a, b = tabelas
        write_comparison_artifacts(resultado, a, b, tmp_path)

        somente_a = pd.read_csv(tmp_path / ONLY_A_CSV_NAME, dtype=str)
        assert list(somente_a.columns) == [*COLUNAS, ROW_COLUMN]
        assert somente_a.iloc[0]["codigo"] == "3"
        assert somente_a.iloc[0][ROW_COLUMN] == "4"

        somente_b = pd.read_csv(tmp_path / ONLY_B_CSV_NAME, dtype=str)
        assert somente_b.iloc[0]["codigo"] == "5"

    def test_csv_vazio_ainda_e_criado_com_cabecalho(self, tmp_path: Path) -> None:
        a = _tabela("a.xlsx", [["1", "Ana", "10,00"]])
        b = _tabela("b.xlsx", [["1", "Ana", "11,00"]])
        resultado = compare_tables(a, b, key_columns=["codigo"])

        write_comparison_artifacts(resultado, a, b, tmp_path)

        conteudo = (tmp_path / ONLY_A_CSV_NAME).read_text(encoding="utf-8-sig")
        assert conteudo.strip() == f"{','.join(COLUNAS)},{ROW_COLUMN}"

    def test_coluna_linha_nao_sobrescreve_coluna_homonima(self, tmp_path: Path) -> None:
        colunas = ["codigo", "linha"]
        a = _tabela("a.xlsx", [["1", "valor original"]], colunas)
        b = _tabela("b.xlsx", [["9", "outro"]], colunas)
        resultado = compare_tables(a, b, key_columns=["codigo"])

        write_comparison_artifacts(resultado, a, b, tmp_path)

        somente_a = pd.read_csv(tmp_path / ONLY_A_CSV_NAME, dtype=str)
        assert somente_a.iloc[0]["linha"] == "valor original"
        assert "linha_" in somente_a.columns


class TestPlanilhaDeDivergencias:
    def test_abas_por_categoria(
        self,
        resultado: ComparisonResult,
        tabelas: tuple[TableSource, TableSource],
        tmp_path: Path,
    ) -> None:
        a, b = tabelas
        write_comparison_artifacts(resultado, a, b, tmp_path)

        wb = load_workbook(tmp_path / XLSX_NAME)
        assert wb.sheetnames == [
            "Resumo",
            "Divergencias",
            "Somente em A",
            "Somente em B",
            "Conflitos de chave",
            "Diferencas toleradas",
        ]

    def test_uma_linha_por_celula_divergente(
        self,
        resultado: ComparisonResult,
        tabelas: tuple[TableSource, TableSource],
        tmp_path: Path,
    ) -> None:
        a, b = tabelas
        write_comparison_artifacts(resultado, a, b, tmp_path)

        ws = load_workbook(tmp_path / XLSX_NAME)["Divergencias"]
        linhas = list(ws.iter_rows(values_only=True))
        assert linhas[0] == (
            "codigo",
            "coluna divergente",
            "valor em A",
            "valor em B",
            "linha em A",
            "linha em B",
        )
        assert linhas[1] == ("2", "valor", "20,00", "25,00", 3, 3)

    def test_resumo_declara_limitacoes(
        self,
        resultado: ComparisonResult,
        tabelas: tuple[TableSource, TableSource],
        tmp_path: Path,
    ) -> None:
        a, b = tabelas
        write_comparison_artifacts(resultado, a, b, tmp_path)

        ws = load_workbook(tmp_path / XLSX_NAME)["Resumo"]
        texto = "\n".join(
            str(valor) for linha in ws.iter_rows(values_only=True) for valor in linha if valor
        )
        assert "Limitacoes desta comparacao" in texto
        assert "Conflitos de chave" in texto

    def test_abas_vazias_nao_quebram(self, tmp_path: Path) -> None:
        a = _tabela("a.xlsx", [["1", "Ana", "10,00"]])
        b = _tabela("b.xlsx", [["1", "Ana", "10,00"]])
        resultado = compare_tables(a, b, key_columns=["codigo"])

        write_comparison_artifacts(resultado, a, b, tmp_path)

        wb = load_workbook(tmp_path / XLSX_NAME)
        assert wb["Divergencias"].max_row == 1  # so o cabecalho


class TestResumoTextual:
    def test_mostra_contadores_e_divergencias(self, resultado: ComparisonResult) -> None:
        texto = generate_summary(resultado)

        assert "Comparacao por chave: codigo" in texto
        assert "Divergentes" in texto
        assert "'20,00'  ->  '25,00'" in texto
        assert "Limitacoes desta comparacao" in texto

    def test_trunca_listas_longas(self) -> None:
        a = _tabela("a.xlsx", [[str(n), f"Nome {n}", "1"] for n in range(10)])
        b = _tabela("b.xlsx", [["999", "Outro", "1"]])
        resultado = compare_tables(a, b, key_columns=["codigo"])

        texto = generate_summary(resultado, max_rows=3)
        assert "... e mais 7" in texto

    def test_conflitos_aparecem_com_as_linhas(self, resultado: ComparisonResult) -> None:
        texto = generate_summary(resultado)
        assert "Conflitos de chave (1)" in texto
        assert "A: 5, 6" in texto
