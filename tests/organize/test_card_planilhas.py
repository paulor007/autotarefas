"""
Card 01 — Análise e organização de planilhas: o contrato, em testes.

Cada classe abaixo corresponde a um item dos critérios de aceite. As
fixtures são de domínios diferentes de propósito (vendas, financeiro,
serviço público, estoque, clientes, pesquisa, atendimentos, contratos):
o card não pode ter sido calibrado para uma planilha só.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from autotarefas.organize import (
    ANALYSIS_REPORT_NAME,
    DASHBOARD_SHEET,
    IndicatorRequest,
    ReportInput,
    SortRequest,
    audit_presentation,
    build_indicators,
    candidates,
    date_format_notes,
    needs_sheet_choice,
    organize_workbook,
    suggest_roles,
    summary_is_offerable,
    survey_sheets,
    text_number_notes,
    write_analysis_report,
)
from autotarefas.reader import read_workbook

FIXTURES = Path(__file__).parent.parent / "fixtures" / "dominios"

VENDAS = FIXTURES / "vendas_simples.xlsx"
FINANCEIRO = FIXTURES / "financeiro_profissional.xlsx"
PUBLICO = FIXTURES / "servico_publico.xlsx"
ESTOQUE = FIXTURES / "estoque_com_formulas.xlsx"
CLIENTES = FIXTURES / "clientes.csv"
PESQUISA = FIXTURES / "pesquisa_ambigua.xlsx"
ATENDIMENTOS = FIXTURES / "atendimentos_duas_abas.xlsx"
CONTRATOS = FIXTURES / "contratos_sem_anomalia.xlsx"


# ============================================================
# Avaliação objetiva da apresentação
# ============================================================


class TestAvaliacaoDaApresentacao:
    def test_planilha_profissional_e_declarada_organizada(self) -> None:
        """Quem já cuidou da planilha não recebe proposta de reformatação."""
        audit = audit_presentation(FINANCEIRO)
        assert audit.verdict == "organizada"
        assert audit.can_improve is False
        assert audit.failures == ()

    def test_planilha_crua_pode_ser_melhorada(self) -> None:
        audit = audit_presentation(VENDAS)
        assert audit.verdict == "melhoravel"
        assert audit.can_improve is True
        pendencias = {c.key for c in audit.failures}
        assert {"cabecalho_destacado", "filtro", "painel_congelado"} <= pendencias

    def test_estrutura_com_mesclagem_e_ambigua(self) -> None:
        """Mesclagem na área de dados: o card não decide sozinho."""
        audit = audit_presentation(PESQUISA)
        assert audit.verdict == "ambigua"
        assert audit.can_improve is False

    def test_cada_criterio_traz_a_evidencia(self) -> None:
        audit = audit_presentation(VENDAS)
        assert all(c.detail for c in audit.criteria)
        assert {c.key for c in audit.criteria} >= {
            "cabecalho_presente",
            "estrutura_tabular",
            "larguras_legiveis",
            "formatos_consistentes",
            "alinhamento_coerente",
            "filtro",
            "painel_congelado",
            "cores_moderadas",
        }

    def test_largura_e_titulo_cortado_num_criterio_so(self) -> None:
        """
        Eram dois criterios e o relatorio se contradizia: "Larguras legiveis:
        OK" seguido de "Titulos visiveis por inteiro: melhorar". Para quem le,
        e a mesma pergunta.
        """
        chaves = {c.key for c in audit_presentation(VENDAS).criteria}
        assert "larguras_legiveis" in chaves
        assert "larguras_adequadas" not in chaves
        assert "texto_nao_cortado" not in chaves

    def test_veredito_e_serializavel(self) -> None:
        payload = audit_presentation(CONTRATOS).as_dict()
        assert payload["veredito"] == "organizada"
        assert payload["criterios"]
        assert payload["pendencias"] == []

    def test_avaliacao_nao_altera_o_arquivo(self) -> None:
        antes = hashlib.sha256(VENDAS.read_bytes()).hexdigest()
        audit_presentation(VENDAS)
        assert hashlib.sha256(VENDAS.read_bytes()).hexdigest() == antes


# ============================================================
# Abas
# ============================================================


class TestAbas:
    def test_classifica_todas_as_abas(self) -> None:
        abas = survey_sheets(ATENDIMENTOS)
        naturezas = {info.name: info.kind for info in abas}
        assert naturezas == {
            "Janeiro": "dados",
            "Fevereiro": "dados",
            "Leia-me": "apresentacao",
            "Rascunho": "vazia",
        }

    def test_duas_candidatas_exigem_escolha(self) -> None:
        abas = survey_sheets(ATENDIMENTOS)
        assert needs_sheet_choice(abas) is True
        assert {c.name for c in candidates(abas)} == {"Janeiro", "Fevereiro"}

    def test_uma_candidata_nao_pergunta(self) -> None:
        assert needs_sheet_choice(survey_sheets(FINANCEIRO)) is False

    def test_cada_aba_explica_a_classificacao(self) -> None:
        assert all(info.reason for info in survey_sheets(ATENDIMENTOS))


# ============================================================
# Organização (formatação opcional e confirmada)
# ============================================================


class TestVariasTabelasNaMesmaAba:
    """
    Duas bases coladas na mesma aba viram uma contagem errada em silencio.

    Nao da para adivinhar qual delas a pessoa quis: o card devolve a pergunta
    e nao oferece organizacao — organizar duas tabelas como se fossem uma so
    seria justamente o chute que o produto evita.
    """

    def _com_duas_tabelas(self, destino: Path) -> Path:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["Aluno", "Turma", "Nota"])
        for i in range(1, 6):
            ws.append([f"Aluno {i}", "3A", 7 + i % 3])
        ws.append([])
        ws.append(["Professor", "Disciplina"])
        for i in range(1, 4):
            ws.append([f"Prof {i}", "Matematica"])
        wb.save(destino)
        return destino

    def test_segunda_tabela_torna_a_estrutura_ambigua(self, tmp_path: Path) -> None:
        audit = audit_presentation(self._com_duas_tabelas(tmp_path / "duas.xlsx"))
        criterio = next(c for c in audit.criteria if c.key == "tabela_unica")

        assert criterio.passed is False
        assert "outra tabela na mesma aba" in criterio.detail
        assert audit.verdict == "ambigua"
        assert audit.can_improve is False

    def test_a_pendencia_diz_o_motivo_e_a_linha(self, tmp_path: Path) -> None:
        """Na tela, o titulo do criterio sozinho nao explica nada."""
        audit = audit_presentation(self._com_duas_tabelas(tmp_path / "duas.xlsx"))
        pendencia = next(
            p for p in audit.as_dict()["pendencias"] if p.startswith("Uma tabela por aba")
        )
        assert "linha 8" in pendencia

    @pytest.mark.parametrize(
        "fixture", [VENDAS, FINANCEIRO, PUBLICO, ESTOQUE, CONTRATOS, ATENDIMENTOS]
    )
    def test_planilha_normal_nao_vira_falso_positivo(self, fixture: Path) -> None:
        criterio = next(c for c in audit_presentation(fixture).criteria if c.key == "tabela_unica")
        assert criterio.passed is True, criterio.detail

    def test_linha_em_branco_no_meio_nao_e_alarme(self, tmp_path: Path) -> None:
        """Base com um respiro no meio e comum demais para virar alarme."""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["Protocolo", "Setor", "Dias"])
        for i in range(1, 5):
            ws.append([f"P{i}", "Obras", i])
        ws.append([])
        for i in range(5, 9):
            ws.append([f"P{i}", "Saude", i])
        destino = tmp_path / "respiro.xlsx"
        wb.save(destino)

        criterio = next(c for c in audit_presentation(destino).criteria if c.key == "tabela_unica")
        assert criterio.passed is True, criterio.detail


class TestDesempenhoDaAvaliacao:
    """
    Regressao de LENTIDAO, que e um defeito como outro qualquer.

    A primeira versao do criterio "uma tabela por aba" pedia a linha inteira ao
    openpyxl uma por vez. Numa planilha de 7 mil linhas isso custava 16
    segundos, e a jornada inteira parecia travada — duas vezes, porque a
    avaliacao roda na analise e de novo no pos-processamento.
    """

    def test_planilha_grande_e_avaliada_em_segundos(self, tmp_path: Path) -> None:
        import time

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["Codigo", "Setor", "Valor"])
        for i in range(5000):
            ws.append([f"P{i}", f"Setor {i % 20}", i * 3])
        destino = tmp_path / "grande.xlsx"
        wb.save(destino)

        inicio = time.perf_counter()
        audit = audit_presentation(destino)
        decorrido = time.perf_counter() - inicio

        assert audit.verdict in {"organizada", "melhoravel", "ambigua"}
        # A implementacao correta leva menos de 1 s; o teto largo e para nao
        # falhar em maquina lenta, e ainda assim pega a regressao de 16 s.
        assert decorrido < 5, f"avaliacao levou {decorrido:.1f}s"


class TestNumeroComoTexto:
    """O leitor entende o valor; dentro do Excel a coluna continua quebrada."""

    def _planilha(self, destino: Path, valores: list[object], formato: str = "@") -> Path:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["Convenio", "Repasse"])
        for i, valor in enumerate(valores, start=2):
            ws.cell(row=i, column=1, value=f"CV-{i:03d}")
            ws.cell(row=i, column=2, value=valor).number_format = formato
        wb.save(destino)
        return destino

    def test_moeda_como_texto_vira_observacao(self, tmp_path: Path) -> None:
        notas = text_number_notes(
            self._planilha(tmp_path / "t.xlsx", ["1234,50", "2500,00", "980,25"])
        )
        assert len(notas) == 1
        assert "Repasse" in notas[0]
        assert "não somam" in notas[0]
        assert "Nada foi convertido" in notas[0]

    @pytest.mark.parametrize("valores", [["15%", "20%", "5%"], ["R$ 10,00", "R$ 20,00"]])
    def test_percentual_e_moeda_escritos_por_gente(
        self, valores: list[object], tmp_path: Path
    ) -> None:
        assert text_number_notes(self._planilha(tmp_path / "p.xlsx", valores))

    def test_numero_de_verdade_nao_gera_ruido(self, tmp_path: Path) -> None:
        assert text_number_notes(self._planilha(tmp_path / "n.xlsx", [1, 2, 3], "General")) == ()

    def test_identificador_com_zero_a_esquerda_nao_e_numero(self, tmp_path: Path) -> None:
        """ "000123" e codigo. Converter seria o estrago que o card evita."""
        assert text_number_notes(self._planilha(tmp_path / "id.xlsx", ["000123", "000124"])) == ()

    def test_texto_comum_nao_e_numero(self, tmp_path: Path) -> None:
        assert (
            text_number_notes(self._planilha(tmp_path / "c.xlsx", ["deferido", "em analise"])) == ()
        )

    def test_nao_altera_o_arquivo(self, tmp_path: Path) -> None:
        arquivo = self._planilha(tmp_path / "t.xlsx", ["1234,50", "2500,00"])
        antes = hashlib.sha256(arquivo.read_bytes()).hexdigest()
        text_number_notes(arquivo)
        assert hashlib.sha256(arquivo.read_bytes()).hexdigest() == antes


class TestAbaGrande:
    """
    Regressao: toda tabela de tamanho REAL era classificada como "ambigua".

    A contagem de celulas olha so as primeiras 50 linhas (amostra), mas a
    densidade era dividida pela area da planilha INTEIRA. Resultado: acima de
    ~111 linhas, qualquer tabela cheia "parecia" 1% preenchida. As fixtures
    tem 13 a 21 linhas, entao a suite inteira passava — foi preciso uma
    planilha de verdade para o furo aparecer.
    """

    def _planilha(self, destino: Path, linhas: int) -> Path:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["Codigo", "Produto", "Valor"])
        for i in range(1, linhas + 1):
            ws.append([i, f"Item {i}", i * 10])
        wb.save(destino)
        return destino

    @pytest.mark.parametrize("linhas", [20, 112, 500, 3000])
    def test_tabela_cheia_e_dados_em_qualquer_tamanho(self, linhas: int, tmp_path: Path) -> None:
        abas = survey_sheets(self._planilha(tmp_path / f"t{linhas}.xlsx", linhas))
        assert abas[0].kind == "dados", abas[0].reason
        assert abas[0].is_candidate is True

    def test_duas_abas_grandes_ainda_pedem_escolha(self, tmp_path: Path) -> None:
        from openpyxl import Workbook

        wb = Workbook()
        for nome in ("Janeiro", "Fevereiro"):
            ws = wb.create_sheet(nome)
            ws.append(["Codigo", "Produto", "Valor"])
            for i in range(1, 400):
                ws.append([i, f"Item {i}", i * 10])
        del wb[wb.sheetnames[0]]
        destino = tmp_path / "duas_grandes.xlsx"
        wb.save(destino)

        abas = survey_sheets(destino)
        assert len(candidates(abas)) == 2
        assert needs_sheet_choice(abas) is True


class TestFormatoDeData:
    """Observar sem alterar: data americana e um deslize que a pessoa tem o direito de saber."""

    def _planilha(self, destino: Path, formato: str) -> Path:
        import datetime as dt

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["Emissao", "Valor"])
        for dia in (1, 2, 3):
            ws.cell(row=dia + 1, column=1, value=dt.date(2019, 12, dia)).number_format = formato
            ws.cell(row=dia + 1, column=2, value=dia * 10)
        wb.save(destino)
        return destino

    def test_data_americana_vira_observacao(self, tmp_path: Path) -> None:
        notas = date_format_notes(self._planilha(tmp_path / "us.xlsx", "mm-dd-yy"))
        assert len(notas) == 1
        assert "Emissao" in notas[0]
        assert "formato americano" in notas[0]
        assert "não foram alterados" in notas[0]

    @pytest.mark.parametrize("formato", ["dd/mm/yyyy", "yyyy-mm-dd", "General"])
    def test_formato_sem_ambiguidade_nao_gera_ruido(self, formato: str, tmp_path: Path) -> None:
        arquivo = self._planilha(tmp_path / f"{formato[:2]}.xlsx", formato)
        assert date_format_notes(arquivo) == ()

    def test_nao_altera_o_arquivo(self, tmp_path: Path) -> None:
        arquivo = self._planilha(tmp_path / "us.xlsx", "mm-dd-yy")
        antes = hashlib.sha256(arquivo.read_bytes()).hexdigest()
        date_format_notes(arquivo)
        assert hashlib.sha256(arquivo.read_bytes()).hexdigest() == antes


class TestDashboard:
    """A aba de painel: opcional, confirmada, e sempre SEPARADA dos dados."""

    PEDIDO = IndicatorRequest(value_column="Valor", category_column="Vendedor")

    def _indicadores(self) -> tuple:
        frame, _ = _frame(VENDAS, "Vendas")
        return build_indicators(frame, self.PEDIDO)

    def test_sem_indicadores_nao_ha_aba(self, tmp_path: Path) -> None:
        destino = tmp_path / "sem.xlsx"
        organize_workbook(VENDAS, destino)
        assert DASHBOARD_SHEET not in load_workbook(destino).sheetnames

    def test_com_papeis_confirmados_a_aba_nasce_na_frente(self, tmp_path: Path) -> None:
        indicadores = self._indicadores()
        assert indicadores, "a fixture precisa render indicadores"
        destino = tmp_path / "com.xlsx"
        resultado = organize_workbook(
            VENDAS,
            destino,
            dashboard=indicadores,
            dashboard_request=self.PEDIDO,
        )
        wb = load_workbook(destino)
        assert wb.sheetnames[0] == DASHBOARD_SHEET
        painel = wb[DASHBOARD_SHEET]
        assert painel.cell(row=1, column=1).value == "Dashboard"
        # A aba declara de onde vieram os numeros — ninguem precisa adivinhar.
        assert "Valor" in str(painel.cell(row=2, column=1).value)
        assert painel._charts, "o painel confirmado tem grafico"
        assert any(m.kind == "dashboard_adicionado" for m in resultado.changes)

    def test_a_aba_de_dados_continua_intocada(self, tmp_path: Path) -> None:
        destino = tmp_path / "com.xlsx"
        organize_workbook(
            VENDAS,
            destino,
            dashboard=self._indicadores(),
            dashboard_request=self.PEDIDO,
        )
        original = load_workbook(VENDAS).active
        dados = load_workbook(destino)["Vendas"]
        assert original is not None
        assert [[c.value for c in linha] for linha in original.iter_rows()] == [
            [c.value for c in linha] for linha in dados.iter_rows()
        ]
        assert not dados._charts, "grafico nenhum entra na aba dos dados"


class TestOrganizacao:
    def test_organizar_resolve_as_pendencias(self, tmp_path: Path) -> None:
        destino = tmp_path / "planilha_organizada.xlsx"
        organize_workbook(VENDAS, destino, sheet="Vendas")
        assert audit_presentation(destino, sheet="Vendas").verdict == "organizada"

    def test_cada_mudanca_e_registrada(self, tmp_path: Path) -> None:
        resultado = organize_workbook(VENDAS, tmp_path / "o.xlsx", sheet="Vendas")
        tipos = {mudanca.kind for mudanca in resultado.changes}
        assert {
            "cabecalho_formatado",
            "largura_ajustada",
            "filtro_aplicado",
            "painel_congelado",
        } <= tipos
        assert all(mudanca.detail for mudanca in resultado.changes)

    def test_valores_nao_mudam_na_organizacao(self, tmp_path: Path) -> None:
        destino = tmp_path / "o.xlsx"
        organize_workbook(VENDAS, destino, sheet="Vendas")

        antes = load_workbook(VENDAS)["Vendas"]
        depois = load_workbook(destino)["Vendas"]
        for linha in range(1, antes.max_row + 1):
            for coluna in range(1, antes.max_column + 1):
                assert (
                    antes.cell(row=linha, column=coluna).value
                    == depois.cell(row=linha, column=coluna).value
                )

    def test_formulas_sobrevivem(self, tmp_path: Path) -> None:
        destino = tmp_path / "o.xlsx"
        organize_workbook(ESTOQUE, destino, sheet="Estoque")
        assert str(load_workbook(destino)["Estoque"]["E2"].value).startswith("=")

    def test_zeros_a_esquerda_sobrevivem(self, tmp_path: Path) -> None:
        destino = tmp_path / "o.xlsx"
        organize_workbook(PUBLICO, destino, sheet="Protocolos")
        assert load_workbook(destino)["Protocolos"]["A2"].value == "000123"

    def test_original_nunca_e_sobrescrito(self, tmp_path: Path) -> None:
        antes = hashlib.sha256(VENDAS.read_bytes()).hexdigest()
        organize_workbook(VENDAS, tmp_path / "o.xlsx", sheet="Vendas")
        assert hashlib.sha256(VENDAS.read_bytes()).hexdigest() == antes

    def test_csv_nao_entra_no_organizador(self, tmp_path: Path) -> None:
        with pytest.raises(ValueError, match="xlsx"):
            organize_workbook(CLIENTES, tmp_path / "o.xlsx")

    def test_correcoes_confirmadas_sao_aplicadas(self, tmp_path: Path) -> None:
        destino = tmp_path / "o.xlsx"
        resultado = organize_workbook(
            PUBLICO,
            destino,
            sheet="Protocolos",
            value_changes=[{"line": 12, "column": "Assunto", "after": "Exame"}],
        )
        assert resultado.value_changes_applied == 1
        assert load_workbook(destino)["Protocolos"]["D12"].value == "Exame"


# ============================================================
# Ordenação
# ============================================================


class TestOrdenacao:
    def test_sem_confirmacao_a_ordem_e_preservada(self, tmp_path: Path) -> None:
        destino = tmp_path / "o.xlsx"
        organize_workbook(PUBLICO, destino, sheet="Protocolos")

        antes = [
            load_workbook(PUBLICO)["Protocolos"].cell(row=linha, column=1).value
            for linha in range(2, 8)
        ]
        depois = [
            load_workbook(destino)["Protocolos"].cell(row=linha, column=1).value
            for linha in range(2, 8)
        ]
        assert antes == depois

    def test_ordenacao_confirmada_e_aplicada(self, tmp_path: Path) -> None:
        destino = tmp_path / "o.xlsx"
        resultado = organize_workbook(
            PUBLICO, destino, sheet="Protocolos", sort=SortRequest("Dias", ascending=False)
        )
        ws = load_workbook(destino)["Protocolos"]
        dias = [ws.cell(row=linha, column=6).value for linha in range(2, 6)]

        assert dias == sorted(dias, reverse=True)
        assert "decrescente" in resultado.sorted_by

    def test_ordenacao_e_recusada_com_formulas(self, tmp_path: Path) -> None:
        """Reordenar quebraria o que cada fórmula aponta."""
        resultado = organize_workbook(
            ESTOQUE, tmp_path / "o.xlsx", sheet="Estoque", sort=SortRequest("SKU")
        )
        assert resultado.sorted_by == ""
        assert any("fórmula" in motivo for motivo in resultado.refusals)

    def test_coluna_inexistente_e_recusada(self, tmp_path: Path) -> None:
        resultado = organize_workbook(
            VENDAS, tmp_path / "o.xlsx", sheet="Vendas", sort=SortRequest("Fantasma")
        )
        assert resultado.sorted_by == ""
        assert any("não existe" in motivo for motivo in resultado.refusals)


# ============================================================
# Indicadores: nada de gráfico inventado
# ============================================================


def _frame(caminho: Path, sheet: str | None = None) -> tuple[pd.DataFrame, list[dict]]:
    leitura = read_workbook(caminho, sheet=sheet)
    assert leitura.original_dataframe is not None
    colunas = [{"name": c.name, "inferred_type": c.inferred_type} for c in leitura.detected_columns]
    return leitura.original_dataframe, colunas


class TestIndicadores:
    def test_sugere_papeis_com_confianca(self) -> None:
        frame, colunas = _frame(VENDAS, "Vendas")
        sugestoes = {s.column: s for s in suggest_roles(frame, colunas)}

        assert sugestoes["Data"].role == "data"
        assert sugestoes["Vendedor"].role == "categoria"
        assert sugestoes["Valor"].role in {"valor", "quantidade"}
        assert all(s.reason for s in sugestoes.values())

    def test_resumo_e_oferecido_quando_ha_dimensao_e_medida(self) -> None:
        frame, colunas = _frame(VENDAS, "Vendas")
        assert summary_is_offerable(suggest_roles(frame, colunas)) is True

    def test_nada_e_calculado_sem_confirmacao(self) -> None:
        """Um pedido vazio não produz número nenhum."""
        frame, _ = _frame(VENDAS, "Vendas")
        assert build_indicators(frame, IndicatorRequest(value_column="")) == ()

    def test_indicador_confirmado_soma_certo(self) -> None:
        frame, _ = _frame(VENDAS, "Vendas")
        pedido = IndicatorRequest(
            value_column="Valor", category_column="Vendedor", date_column="Data"
        )
        indicadores = build_indicators(frame, pedido)

        assert len(indicadores) == 2
        por_vendedor = indicadores[0]
        assert por_vendedor.dimension == "Vendedor"

        # A soma é conferida contra a própria coluna, não contra um número
        # decorado: assim o teste continua válido se a fixture crescer.
        esperado = sum(float(str(v).replace(",", ".")) for v in frame["Valor"])
        assert por_vendedor.total == pytest.approx(esperado)
        assert sum(valor for _, valor in por_vendedor.rows) == pytest.approx(esperado)

    def test_indicador_por_mes_agrupa_por_competencia(self) -> None:
        frame, _ = _frame(VENDAS, "Vendas")
        pedido = IndicatorRequest(value_column="Valor", date_column="Data")
        (por_mes,) = build_indicators(frame, pedido)

        assert por_mes.dimension == "mês"
        assert [chave for chave, _ in por_mes.rows] == ["2026-01", "2026-02", "2026-03"]

    def test_coluna_de_valor_inexistente_nao_inventa(self) -> None:
        frame, _ = _frame(VENDAS, "Vendas")
        pedido = IndicatorRequest(value_column="Fantasma", category_column="Vendedor")
        assert build_indicators(frame, pedido) == ()

    def test_dominio_publico_tambem_e_suportado(self) -> None:
        frame, colunas = _frame(PUBLICO, "Protocolos")
        sugestoes = {s.column: s.role for s in suggest_roles(frame, colunas)}
        assert sugestoes["Orgao"] == "categoria"
        assert sugestoes["Abertura"] == "data"


# ============================================================
# Relatório de análise
# ============================================================


class TestRelatorio:
    def _entrada(self, **extra) -> ReportInput:
        return ReportInput(
            source_name=VENDAS.name,
            sheet="Vendas",
            rows=20,
            columns=6,
            audit=audit_presentation(VENDAS),
            sheets=survey_sheets(VENDAS),
            **extra,
        )

    def test_gera_o_arquivo_com_o_resumo(self, tmp_path: Path) -> None:
        destino = write_analysis_report(tmp_path / ANALYSIS_REPORT_NAME, self._entrada())
        wb = load_workbook(destino)

        assert "Resumo" in wb.sheetnames
        texto = "\n".join(
            str(v) for linha in wb["Resumo"].iter_rows(values_only=True) for v in linha if v
        )
        assert "pode ser melhorada" in texto
        assert VENDAS.name in texto

    def test_abas_vazias_nao_sao_escritas(self, tmp_path: Path) -> None:
        destino = write_analysis_report(tmp_path / "r.xlsx", self._entrada())
        wb = load_workbook(destino)
        assert "Antes e depois" not in wb.sheetnames

    def test_antes_e_depois_registra_valor_anterior_e_posterior(self, tmp_path: Path) -> None:
        entrada = self._entrada(
            cleaning_changes=[
                {
                    "line": 5,
                    "column": "Vendedor",
                    "before": "  Ana  ",
                    "after": "Ana",
                    "rules": ["espacos"],
                }
            ]
        )
        destino = write_analysis_report(tmp_path / "r.xlsx", entrada)
        ws = load_workbook(destino)["Antes e depois"]

        assert [c.value for c in ws[1]] == [
            "linha",
            "coluna",
            "valor anterior",
            "valor posterior",
            "tipo da alteração",
            "motivo",
        ]
        assert [c.value for c in ws[2]][:4] == [5, "Vendedor", "  Ana  ", "Ana"]

    def test_alteracoes_visuais_ficam_registradas(self, tmp_path: Path) -> None:
        resultado = organize_workbook(VENDAS, tmp_path / "o.xlsx", sheet="Vendas")
        destino = write_analysis_report(tmp_path / "r.xlsx", self._entrada(organize=resultado))
        ws = load_workbook(destino)["Alteracoes realizadas"]
        tipos = {linha[1] for linha in ws.iter_rows(min_row=2, values_only=True)}

        assert "filtro_aplicado" in tipos
        assert "cabecalho_formatado" in tipos

    def test_indicadores_so_aparecem_quando_confirmados(self, tmp_path: Path) -> None:
        sem = write_analysis_report(tmp_path / "sem.xlsx", self._entrada())
        assert "Indicadores confirmados" not in load_workbook(sem).sheetnames

        frame, _ = _frame(VENDAS, "Vendas")
        indicadores = build_indicators(
            frame, IndicatorRequest(value_column="Valor", category_column="Vendedor")
        )
        com = write_analysis_report(tmp_path / "com.xlsx", self._entrada(indicators=indicadores))
        wb = load_workbook(com)
        assert "Indicadores confirmados" in wb.sheetnames
        assert wb["Indicadores confirmados"]._charts  # o gráfico foi criado

    def test_abas_do_arquivo_listam_a_classificacao(self, tmp_path: Path) -> None:
        entrada = ReportInput(
            source_name=ATENDIMENTOS.name,
            sheet="Janeiro",
            sheets=survey_sheets(ATENDIMENTOS),
        )
        destino = write_analysis_report(tmp_path / "r.xlsx", entrada)
        ws = load_workbook(destino)["Abas do arquivo"]
        naturezas = {linha[0]: linha[1] for linha in ws.iter_rows(min_row=2, values_only=True)}

        assert naturezas["Rascunho"] == "vazia"
        assert naturezas["Leia-me"] == "apresentacao"


def test_fixtures_de_dominio_existem() -> None:
    for caminho in (
        VENDAS,
        FINANCEIRO,
        PUBLICO,
        ESTOQUE,
        CLIENTES,
        PESQUISA,
        ATENDIMENTOS,
        CONTRATOS,
    ):
        assert caminho.is_file(), caminho
